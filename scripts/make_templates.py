"""Generate sample upload templates from the importer's own field definitions.

Run from the repository root:  python scripts/make_templates.py
Writes templates/<type>_template.xlsx and .csv, and copies them to frontend/public/templates
so the UI can offer them for download. Templates contain headers and format guidance
only: no example prices, so nothing in them can be mistaken for market evidence.
"""
from __future__ import annotations

import csv
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from app.importing.fields import RECORD_TYPES, fields_for  # noqa: E402
from app.pricing.config import default_config  # noqa: E402

_CFG = default_config()
# Columns with a fixed set of accepted values get an Excel dropdown on the Data sheet.
CHOICES = {
    "grade": _CFG.grade_scale,
    "bool": ["yes", "no"],
    "tax_basis": ["inclusive", "exclusive", "unknown"],
    "channel": ["retail", "wholesale", "as_is"],
    "availability": ["in_stock", "out_of_stock", "sold", "removed", "unknown"],
    "txn_kind": ["sale", "purchase"],
    "fault_category": [f.key for f in _CFG.fault_categories],
    "storage_type": ["SSD", "HDD", "SSHD", "eMMC"],
}
DATA_ROWS = 2000

FORMAT_HINTS = {
    "text": "Free text", "cpu": "Exact model, e.g. i5-8350U, Ryzen 5 PRO 4650U, Apple M1",
    "ram": "GB, e.g. 16 or 16GB", "storage": "GB or TB, e.g. 512, 512GB SSD, 1TB",
    "storage_type": "SSD / HDD / SSHD / eMMC", "display_size": "Inches, e.g. 14 or 15.6\"",
    "resolution": "e.g. 1920x1080, FHD, QHD", "grade": "One of the configured grades (default A+, A, B, C, D)",
    "percent": "e.g. 85 or 85%", "pct": "Percent, e.g. 15 or 15%", "faults": "Semicolon-separated; 'none' = no faults; blank = unknown",
    "fault_category": "none / cosmetic / battery / display / keyboard_input / board_power / other",
    "bool": "yes / no", "warranty": "e.g. 6 months, 1 year, none", "quantity": "Whole number >= 1",
    "money": "INR per unit; blank = unknown, 0 = none", "price": "INR per unit as observed (must be > 0)",
    "tax_basis": "inclusive / exclusive / unknown (GST)", "currency": "ISO code, default INR",
    "channel": "retail / wholesale / as_is", "availability": "in_stock / out_of_stock / sold / removed",
    "txn_kind": "sale / purchase", "date": "YYYY-MM-DD or DD/MM/YYYY", "datetime": "YYYY-MM-DD HH:MM",
    "url": "Only a real http(s) URL you observed",
}

HEAD_FILL = PatternFill("solid", fgColor="291F56")  # Casterly indigo
REQ_FILL = PatternFill("solid", fgColor="514A9D")  # Casterly purple: required columns
LOGO = ROOT / "backend" / "app" / "assets" / "casterly-logo.png"


def main() -> None:
    out = ROOT / "templates"
    out.mkdir(exist_ok=True)
    for rt, label in RECORD_TYPES.items():
        fields = fields_for(rt)
        headers = [f.label for f in fields]
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(headers)
        for i, (c, f) in enumerate(zip(ws[1], fields), start=1):
            required = rt in f.required
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = REQ_FILL if required else HEAD_FILL
            hint = FORMAT_HINTS.get(f.kind, "")
            note = ("REQUIRED. " if required else "") + hint + (f". {f.help}" if f.help else "")
            c.comment = Comment(note.strip(), "Casterly")
            col = get_column_letter(i)
            ws.column_dimensions[col].width = max(14, len(f.label) + 4)
            choices = CHOICES.get(f.kind)
            if choices:
                dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True,
                                    showErrorMessage=True, errorStyle="warning",
                                    errorTitle=f.label, error="Not one of the usual values: " + ", ".join(choices),
                                    showInputMessage=True, promptTitle=f.label[:32], prompt=hint[:255])
                dv.add(f"{col}2:{col}{DATA_ROWS + 1}")
                ws.add_data_validation(dv)
            if f.kind in ("money", "price"):
                for r in range(2, DATA_ROWS + 2):
                    ws.cell(r, i).number_format = "#,##0.00"
            elif f.kind in ("date",):
                for r in range(2, DATA_ROWS + 2):
                    ws.cell(r, i).number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
        info = wb.create_sheet("Instructions")
        info.append([f"{label}: column guide"])
        info["A1"].font = Font(bold=True, size=13)
        info.append(["Fill the 'Data' sheet, one row per line, keeping the header row. Purple headers are required. "
                     "Hover a header for its format; dropdowns show accepted values. Then upload this file in "
                     "Upload & review."])
        info.append(["Leave a cell blank when the value is unknown. Enter 0 only when the value really is zero."])
        info.append([])
        info.append(["Column", "Required", "Needed for exact matching", "Format", "Notes"])
        for c in info[5]:
            c.font = Font(bold=True)
        for f in fields:
            info.append([f.label, "yes" if rt in f.required else "", "yes" if f.matching else "",
                         FORMAT_HINTS.get(f.kind, ""), f.help])
        for col, w in zip("ABCDE", (30, 10, 24, 60, 60)):
            info.column_dimensions[col].width = w
        logo = XLImage(str(LOGO))
        ratio = logo.width / logo.height
        logo.height, logo.width = 96, round(96 * ratio)
        info.add_image(logo, "G2")
        wb.save(out / f"{rt}_template.xlsx")
        with open(out / f"{rt}_template.csv", "w", newline="", encoding="utf-8-sig") as fh:
            csv.writer(fh).writerow(headers)
    pub = ROOT / "frontend" / "public" / "templates"
    pub.mkdir(parents=True, exist_ok=True)
    for p in out.glob("*_template.*"):
        shutil.copy(p, pub / p.name)
    print(f"Wrote {len(list(out.glob('*_template.*')))} templates to {out} and {pub}")


if __name__ == "__main__":
    main()
