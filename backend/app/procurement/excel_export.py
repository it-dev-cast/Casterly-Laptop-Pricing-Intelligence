"""Priced workbook export.

The vendor's original workbook (as uploaded, never modified on disk) is re-opened with formatting
and formulas intact. The ``Market Unit Price`` column is filled where it already exists (B002
column J) or appended after the last header (B001 and other layouts, with an ``Estimated Lot Value``
formula when a quantity column exists). Supplementary sheets are added; an existing sheet of the
same name is never overwritten (ours are suffixed "(CLPI)"). Every web- or vendor-sourced string
is neutralised against formula injection. The saved file is re-read to verify that the number of
filled prices equals the number of priced assets.
"""
from __future__ import annotations

import io
from collections import defaultdict
from copy import copy
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..importing.files import stored_path
from ..services.export import HEADER_FILL, HEADER_FONT, add_logo, safe
from .excel_import import norm_header
from .models import (PIAsset, PIBatch, PIEstimate, PIException, PIObservation, PIProduct, PIQuotation,
                     PIQuotationLine, PIRuleVersion, PISource)
from .rules import PRICE_TYPES
from .workbook_style import asset_style, style_table, NAVY

INR_FMT = '"₹" #,##0;[Red]-"₹" #,##0'
TITLE_FONT = Font(bold=True, size=14, color="291F56")
NOTE_FILL = PatternFill("solid", fgColor="F3F8DD")
SHEETS = ("Valuation Summary", "Model Pricing Evidence", "Review Before Purchase", "Rate Benchmarks", "Research Sources")


class ExportIntegrityError(RuntimeError):
    pass


def _sheet(wb, title: str):
    name = title
    if name in wb.sheetnames:
        name = f"{title} (CLPI)"[:31]
        suffix = 2
        while name in wb.sheetnames:
            name = f"{title[:22]} (CLPI {suffix})"[:31]
            suffix += 1
    return wb.create_sheet(name)


def _table(ws, start_row: int, header: list[str], rows: list[list], money_cols: set[int] = frozenset(),
           widths: dict[int, int] | None = None) -> int:
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            if isinstance(v, Decimal):
                v = float(v)
            c = ws.cell(row=i, column=j, value=safe(v))
            if j in money_cols:
                c.number_format = INR_FMT
    for j in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 16)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    style_table(ws, start_row, start_row + len(rows), len(header))
    return start_row + len(rows)


def build(db: Session, b: PIBatch, quotation: PIQuotation | None = None, *, preserve: bool = True) -> tuple[bytes, dict]:
    if preserve:
        try:
            wb = load_workbook(stored_path(b.stored_name))
        except Exception as exc:
            raise ExportIntegrityError("Original workbook cannot be safely processed; use the alternate clean quotation export") from exc
    else:
        from openpyxl import Workbook
        wb = Workbook()
        clean = wb.active
        clean.title = b.sheet_name
        source_headers = (b.detection or {}).get("headers", [])
        for col, header in enumerate(source_headers, 1):
            clean.cell(row=b.header_row or 1, column=col, value=safe(header))
        for asset in db.scalars(select(PIAsset).where(PIAsset.batch_id == b.id)).all():
            for col, header in enumerate(source_headers, 1):
                clean.cell(row=asset.row_number, column=col, value=safe(asset.raw.get(header)))
    ws = wb[b.sheet_name]
    hr = b.header_row or 1
    headers = {norm_header(ws.cell(row=hr, column=c).value): c for c in range(1, ws.max_column + 1)
               if ws.cell(row=hr, column=c).value is not None}
    last_col = max(headers.values()) if headers else ws.max_column
    pcol = headers.get(norm_header(b.price_column)) if b.price_column else headers.get("market unit price")
    created_cols = []
    if pcol is None:
        pcol = last_col + 1
        src = ws.cell(row=hr, column=last_col)
        c = ws.cell(row=hr, column=pcol, value="Market Unit Price")
        c.font, c.fill, c.border, c.alignment = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment)
        ws.column_dimensions[get_column_letter(pcol)].width = 18
        created_cols.append("Market Unit Price")
    qcol = headers.get(norm_header(b.qty_column)) if b.qty_column else None
    vcol = None
    if qcol:
        vcol = headers.get("estimated lot value")
        if vcol is None:
            vcol = max(pcol, last_col) + 1
            src = ws.cell(row=hr, column=pcol)
            c = ws.cell(row=hr, column=vcol, value="Estimated Lot Value")
            c.font, c.fill, c.border, c.alignment = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment)
            ws.column_dimensions[get_column_letter(vcol)].width = 20
            created_cols.append("Estimated Lot Value")

    assets = db.scalars(select(PIAsset).where(PIAsset.batch_id == b.id).order_by(PIAsset.row_number)).all()
    qlines = {}
    if quotation is not None:
        qlines = {ln.asset_id: ln for ln in db.scalars(select(PIQuotationLine).where(
            PIQuotationLine.quotation_id == quotation.id)).all()}
    filled = 0
    initial_total = Decimal("0")
    cached_line_values = {}
    first_row = min((a.row_number for a in assets), default=hr + 1)
    last_row = max((a.row_number for a in assets), default=hr + 1)
    pl, ql = get_column_letter(pcol), get_column_letter(qcol) if qcol else None
    for a in assets:
        price = (qlines[a.id].unit_price if a.id in qlines else None) if quotation is not None else a.final_price
        cell = ws.cell(row=a.row_number, column=pcol)
        cell.value = float(price) if price is not None else None
        if price is not None:
            filled += 1
            line_total = Decimal(str(price)) * Decimal(str(a.quantity or 0)) if qcol else Decimal(str(price))
            initial_total += line_total
            if vcol:
                cached_line_values[f"{get_column_letter(vcol)}{a.row_number}"] = line_total
        if not cell.number_format or cell.number_format == "General":
            cell.number_format = INR_FMT
        if vcol:
            vc = ws.cell(row=a.row_number, column=vcol, value=(f"=IF(ISNUMBER({pl}{a.row_number}),"
                                                                     f"N({ql}{a.row_number})*{pl}{a.row_number},\"\")"))
            vc.number_format = INR_FMT

    sname = quote_sheetname(ws.title)
    footer_rows, total_cells = asset_style(ws, hr, first_row, last_row, pcol, vcol, len(assets) - filled)
    price_rng = f"{sname}!${pl}${first_row}:${pl}${last_row}"
    total_formula = (f"=SUMPRODUCT({sname}!${ql}${first_row}:${ql}${last_row},{price_rng})" if ql
                     else f"=SUM({price_rng})")

    products = {p.id: p for p in db.scalars(select(PIProduct).where(PIProduct.id.in_({a.product_id for a in assets if a.product_id}))).all()}
    estimates = {e.id: e for e in db.scalars(select(PIEstimate).where(PIEstimate.batch_id == b.id,
                                                                     PIEstimate.is_current.is_(True))).all()}
    rv = db.get(PIRuleVersion, b.rule_version_id) if b.rule_version_id else None

    # ---------------------------------------------------------------- Model Pricing Evidence
    groups: dict[int | None, list[PIAsset]] = defaultdict(list)
    for a in assets:
        groups[a.estimate_id].append(a)
    mws = _sheet(wb, "Model Pricing Evidence")
    mrows = []
    for eid, members in sorted(groups.items(), key=lambda kv: (kv[0] is None, kv[0] or 0)):
        e = estimates.get(eid) if eid else None
        units = sum(a.quantity or 0 for a in members)
        finals = {(qlines[a.id].unit_price if a.id in qlines else None) if quotation is not None else a.final_price
                  for a in members}
        uniform = len(finals) == 1
        unit = next(iter(finals)) if uniform else None
        p = products.get(members[0].product_id)
        br = e.breakdown if e else {}
        used = [o for o in br.get("observations", []) if o.get("included")]
        urls = [o["url"] for o in used if o.get("url")][:3]
        mrows.append([CATS.get(members[0].category, members[0].category), p.label if p else (members[0].original_spec or "Unidentified"),
                      members[0].grade or "", len(members), units, e.benchmark if e else None,
                      unit if uniform else "mixed (overrides)", None, e.confidence if e else None,
                      (e.status if e else "manual_required"), len(used),
                      ", ".join(sorted({PRICE_TYPES.get(o["price_type"], o["price_type"]) for o in used})),
                      " | ".join(urls), br.get("message")])
    end = _table(mws, 1, ["Category", "Normalised product", "Grade", "Lines", "Units", "Market benchmark (₹)",
                          "Market Unit Price (₹)", "Lot value (₹)", "Confidence", "Status", "Evidence used",
                          "Evidence types", "Source URLs", "Basis"], mrows, {6, 7, 8},
                 {1: 16, 2: 40, 7: 18, 8: 18, 12: 30, 13: 60, 14: 40})
    for i in range(2, end + 1):
        mws.cell(row=i, column=8, value=f"=IF(ISNUMBER(G{i}),E{i}*G{i},\"\")").number_format = INR_FMT

    # ---------------------------------------------------------------- Valuation Summary
    sws = _sheet(wb, "Valuation Summary")
    add_logo(sws, "E1", 140)
    sws["A1"] = "Casterly Asset Valuation"
    sws["A1"].font = Font(name="Calibri", bold=True, size=15, color=NAVY)
    sws.merge_cells("A1:D1")
    sws.row_dimensions[1].height = 34
    sws.sheet_view.showGridLines = False
    sws.sheet_properties.tabColor = NAVY
    meta = [("Batch", safe(b.name)), ("Vendor", safe(quotation.vendor if quotation else b.vendor) or "—"),
            ("Source file", safe(b.original_filename)), ("Asset sheet", safe(b.sheet_name)),
            ("Detected layout", b.layout.upper()),
            ("Priced at", b.priced_at.strftime("%d %b %Y %H:%M UTC") if b.priced_at else "not priced"),
            ("Rule version", f"v{rv.version} · {safe(rv.name)}" if rv else "—"),
            ("Currency / basis", "INR per unit, excluding GST"),
            ("Price definition", "Recommended Casterly vendor purchase price (buy-side); not a retail or resale price"),
            ("Quotation", f"{quotation.quote_ref} (valid until {quotation.valid_until:%d %b %Y})" if quotation else "—")]
    for i, (k, v) in enumerate(meta, start=3):
        sws.cell(row=i, column=1, value=k).font = Font(bold=True)
        sws.cell(row=i, column=2, value=v)
        sws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
        sws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="center")
        sws.row_dimensions[i].height = 30 if k == "Price definition" else 24
    r0 = 3 + len(meta) + 1
    cats = sorted({row[0] for row in mrows})
    me = quote_sheetname(mws.title)
    srows = []
    for c in cats:
        srows.append([c, f"=SUMIF({me}!$A:$A,A{{r}},{me}!$D:$D)", f"=SUMIF({me}!$A:$A,A{{r}},{me}!$E:$E)",
                      f"=SUMIF({me}!$A:$A,A{{r}},{me}!$H:$H)"])
    hdr = ["Category", "Lines", "Units", "Estimated purchase value (₹)"]
    for j, h in enumerate(hdr, start=1):
        cc = sws.cell(row=r0, column=j, value=h)
        cc.fill, cc.font = PatternFill("solid", fgColor=NAVY), HEADER_FONT
    for i, row in enumerate(srows, start=r0 + 1):
        for j, v in enumerate(row, start=1):
            sws.cell(row=i, column=j, value=v.replace("{r}", str(i)) if isinstance(v, str) and v.startswith("=") else safe(v))
        sws.cell(row=i, column=4).number_format = INR_FMT
    tr = r0 + len(srows) + 1
    sws.cell(row=tr, column=1, value="ALL ASSETS").font = Font(bold=True)
    for j, col in ((2, "B"), (3, "C"), (4, "D")):
        cc = sws.cell(row=tr, column=j, value=f"=SUM({col}{r0 + 1}:{col}{tr - 1})")
        cc.font = Font(bold=True)
    sws.cell(row=tr, column=4).number_format = INR_FMT
    for cc in sws[tr][:4]:
        cc.fill = PatternFill("solid", fgColor=NAVY)
        cc.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    sws.row_dimensions[tr].height = 28
    chk = tr + 2
    checks = [("Asset-sheet total (formula over Market Unit Price)", total_formula, True),
              ("Difference vs category total (should be 0)", f"=ROUND(B{chk}-D{tr},2)", True),
              ("Rows imported", len(assets), False), ("Rows with a Market Unit Price", filled, False),
              ("Rows needing manual pricing (blank price)", len(assets) - filled, False)]
    for i, (k, v, is_money) in enumerate(checks, start=chk):
        sws.cell(row=i, column=1, value=k)
        cc = sws.cell(row=i, column=2, value=v)
        if is_money:
            cc.number_format = INR_FMT
    notes_r = chk + len(checks) + 1
    for i, t in enumerate([
        "How to use: Market Unit Price is the recommended maximum starting buy price for the stated condition, ex-GST.",
        "Blank prices need manual pricing: see Review Before Purchase. Provisional prices rest only on historical or "
        "comparable evidence.",
        "Evidence and calculation per product are on Model Pricing Evidence and Research Sources; rules on Rate Benchmarks.",
        "Confirm quantities, condition, testing, warranty/returns, GST and logistics with the vendor before ordering."],
            start=notes_r):
        sws.cell(row=i, column=1, value=t).fill = NOTE_FILL
    sws.column_dimensions["A"].width = 52
    sws.column_dimensions["B"].width = 44
    sws.column_dimensions["C"].width = 14
    sws.column_dimensions["D"].width = 28
    sws.column_dimensions["E"].width = 23
    for row in sws.iter_rows(min_row=3, max_row=notes_r+3, max_col=5):
        for cell in row:
            ff = copy(cell.font)
            ff.name = "Calibri"
            cell.font = ff
    for rr in range(notes_r, notes_r+4):
        sws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=5)
        sws.cell(rr, 1).alignment = Alignment(wrap_text=True, vertical="center")
        sws.row_dimensions[rr].height = 30

    # ---------------------------------------------------------------- Review Before Purchase
    rws = _sheet(wb, "Review Before Purchase")
    row_of = {a.id: a for a in assets}
    exc = db.scalars(select(PIException).where(PIException.batch_id == b.id).order_by(PIException.severity,
                                                                                    PIException.code)).all()
    sev_rank = {"high": 0, "medium": 1, "low": 2}
    erows = []
    for e in sorted(exc, key=lambda x: (sev_rank.get(x.severity, 3), x.code, x.id)):
        a = row_of.get(e.asset_id) if e.asset_id else None
        p = products.get(e.product_id) if e.product_id else None
        erows.append([e.severity.upper(), e.code, a.row_number if a else "all rows of product", a.serial if a else None,
                      (p.label if p else (a.original_spec if a else None)), e.message, e.impact, e.action,
                      "Resolved" if e.resolved else "Open", e.resolution_note])
    _table(rws, 1, ["Severity", "Issue", "Excel row", "Serial", "Product", "Detail", "Business impact",
                    "Verification action", "Status", "Resolution note"], erows, set(),
           {2: 24, 5: 36, 6: 50, 7: 50, 8: 50, 10: 30})

    # ---------------------------------------------------------------- Rate Benchmarks
    bws = _sheet(wb, "Rate Benchmarks")
    brows = []
    for e in sorted(estimates.values(), key=lambda x: (products[x.product_id].category, products[x.product_id].label)):
        p = products[e.product_id]
        br = e.breakdown
        steps = {s["step"]: s for s in br.get("steps", [])}
        brows.append([CATS.get(p.category, p.category), p.label, e.grade, e.group_quantity, e.benchmark,
                      steps.get("condition", {}).get("pct"), steps.get("quantity", {}).get("pct"),
                      steps.get("risk", {}).get("pct"), e.recommended_unit, e.opening_unit, e.max_unit, e.resale_unit,
                      e.margin_pct, e.confidence, e.confidence_score, e.status,
                      "; ".join(br.get("assumptions", []))[:2000]])
    _table(bws, 1, ["Category", "Normalised product", "Grade", "Group units", "Benchmark (₹)", "Grade adj %",
                    "Quantity adj %", "Risk adj %", "Recommended (₹)", "Opening offer (₹)", "Max permitted (₹)",
                    "Expected resale (₹)", "Margin %", "Confidence", "Score", "Status", "Assumptions"], brows,
           {5, 9, 10, 11, 12}, {2: 40, 17: 80})

    # ---------------------------------------------------------------- Research Sources
    rsw = _sheet(wb, "Research Sources")
    used_ids = {o["id"]: o for e in estimates.values() for o in e.breakdown.get("observations", [])}
    srcs = {s.id: s for s in db.scalars(select(PISource)).all()}
    obs = db.scalars(select(PIObservation).where(PIObservation.product_id.in_(list(products) or [0]))
                     .order_by(PIObservation.product_id, PIObservation.id)).all()
    orows = []
    for o in obs:
        u = used_ids.get(o.id)
        s = srcs.get(o.source_id) if o.source_id else None
        orows.append([products[o.product_id].label, PRICE_TYPES.get(o.price_type, o.price_type), o.match_quality,
                      o.price, o.tax_basis, u.get("ex_gst") if u else None, u.get("converted") if u else None,
                      "Used" if u and u.get("included") else (u.get("exclusion") if u else f"not used ({o.status})"),
                      o.status, s.name if s else o.seller, o.url, o.observed_at.strftime("%d %b %Y"), o.provider,
                      (o.title or "")[:300], (o.attributes or {}).get("research_mode", "historical" if o.price_type == "historical_valuation" else "verified_record"),
                      (o.attributes or {}).get("source_timestamp"), o.retrieved_at.isoformat() if o.retrieved_at else None,
                      o.status_reason])
    _table(rsw, 1, ["Product", "Price type", "Match", "Observed price (₹)", "Tax basis", "Ex-GST (₹)",
                    "Buy-side equivalent (₹)", "Use in this valuation", "Verification", "Source", "URL", "Observed",
                    "Provider", "Title", "Research method", "Source timestamp", "Research timestamp", "Validation reason"], orows, {4}, {1: 36, 8: 40, 10: 22, 11: 60, 14: 50})

    if quotation is not None:
        qws = _sheet(wb, "Vendor Quotation")
        qws["A1"] = f"Vendor purchase quotation {quotation.quote_ref}"
        qws["A1"].font = TITLE_FONT
        lines = sorted(qlines.values(), key=lambda ln: row_of[ln.asset_id].row_number)
        qrows = [[row_of[ln.asset_id].row_number, row_of[ln.asset_id].serial, ln.description, ln.quantity, ln.unit_price,
                  None, ln.price_basis] for ln in lines]
        end = _table(qws, 3, ["Excel row", "Serial", "Description", "Qty", "Unit price (₹, ex-GST)", "Line total (₹)",
                              "Basis"], qrows, {5, 6}, {3: 50, 5: 18, 6: 18})
        for i in range(4, end + 1):
            qws.cell(row=i, column=6, value=f"=D{i}*E{i}").number_format = INR_FMT
        qws.cell(row=end + 1, column=5, value="TOTAL").font = Font(bold=True)
        tc = qws.cell(row=end + 1, column=6, value=f"=SUM(F4:F{end})")
        tc.font, tc.number_format = Font(bold=True), INR_FMT

    out = io.BytesIO()
    wb.save(out)
    data = out.getvalue()
    from .xlsx_patch import cache_formula_values
    summary_cache = {f"B{chk}": initial_total}
    model_cache = {}
    category_total = Decimal("0")
    for i, row in enumerate(mrows, start=2):
        if isinstance(row[6], (Decimal, int, float)):
            model_cache[f"H{i}"] = Decimal(str(row[4])) * Decimal(str(row[6]))
    for i, category in enumerate(cats, start=r0+1):
        matching = [(idx, row) for idx, row in enumerate(mrows, start=2) if row[0] == category]
        summary_cache[f"B{i}"] = sum(row[3] for _, row in matching)
        summary_cache[f"C{i}"] = sum(row[4] for _, row in matching)
        value = sum((model_cache.get(f"H{idx}", Decimal("0")) for idx, _ in matching), Decimal("0"))
        summary_cache[f"D{i}"] = value
        category_total += value
    summary_cache.update({f"B{tr}": len(assets), f"C{tr}": sum(a.quantity or 0 for a in assets),
                          f"D{tr}": category_total, f"B{chk+1}": initial_total-category_total})
    data = cache_formula_values(data, {
        ws.title: {**cached_line_values, **{ref: initial_total for ref in total_cells}},
        sws.title: summary_cache,
        mws.title: model_cache,
    })
    fidelity = {"mode": "clean_quotation", "warning": "Clean reserialized workbook; unsupported original Excel objects may not survive. Original upload is unchanged."}
    if preserve:
        from .xlsx_patch import patch
        data, fidelity = patch(stored_path(b.stored_name).read_bytes(), data, b.sheet_name, pcol, vcol,
                               hr, [a.row_number for a in assets], footer_rows=footer_rows)
    integrity = verify(data, b.sheet_name, pcol, [a.row_number for a in assets], filled)
    integrity.update(created_columns=created_cols, price_column=get_column_letter(pcol), fidelity=fidelity)
    return data, integrity


CATS = {"laptop": "Laptop", "mobile_workstation": "Mobile workstation", "macbook": "Apple MacBook",
        "desktop": "Desktop", "smartphone": "Smartphone", "tablet": "Tablet", "monitor": "Monitor",
        "docking_station": "Docking station", "accessory": "Accessory", "server_cpu": "PM-CPU · Server processor",
        "server_ram": "RC-RAM · Server memory", "other": "Other"}


def verify(data: bytes, sheet: str, pcol: int, rows: list[int], expected_filled: int) -> dict:
    wb = load_workbook(io.BytesIO(data), read_only=True)
    try:
        ws = wb[sheet]
        rowset = set(rows)
        filled = 0
        for i, r in enumerate(ws.iter_rows(min_col=pcol, max_col=pcol, values_only=True), start=1):
            if i in rowset and isinstance(r[0], (int, float)):
                filled += 1
        if filled != expected_filled:
            raise ExportIntegrityError(f"Export check failed: {filled} prices written, {expected_filled} expected")
        missing = [s for s in SHEETS if s not in wb.sheetnames and f"{s} (CLPI)"[:31] not in wb.sheetnames]
        if missing:
            raise ExportIntegrityError(f"Export check failed: missing sheets {missing}")
        return {"rows_imported": len(rows), "prices_written": filled, "sheets": wb.sheetnames}
    finally:
        wb.close()
