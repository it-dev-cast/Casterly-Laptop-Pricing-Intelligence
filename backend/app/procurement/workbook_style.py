"""Casterly document-template styling for generated valuation exports."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "0B2545"
BLUE = "1F5AA6"
PALE = "EAF1FB"
LIME = "C4DB61"
MONEY = '"₹" #,##0;[Red]-"₹" #,##0'


def style_table(ws, header_row, end_row, end_col):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = NAVY
    ws.freeze_panes = f"B{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(end_col)}{end_row}"
    ws.row_dimensions[header_row].height = 34
    for row in ws.iter_rows(min_row=header_row, max_row=end_row, max_col=end_col):
        for c in row:
            is_header = c.row == header_row
            c.font = Font(name="Calibri", size=11, color="FFFFFF" if is_header else "404040", bold=is_header)
            c.fill = PatternFill("solid", fgColor=NAVY if is_header else ("F4F7FB" if (c.row-header_row) % 2 == 0 else "FFFFFF"))
            c.alignment = Alignment(vertical="center", horizontal="center" if is_header else ("right" if isinstance(c.value, (float, int)) else "left"), wrap_text=is_header)
            c.border = Border(bottom=Side(style="thin", color="4F86C6")) if is_header else Border()
        if row[0].row != header_row:
            ws.row_dimensions[row[0].row].height = 21
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddFooter.center.text = "Casterly Private Limited | &P of &N"
    ws.oddFooter.center.size = 9


def asset_style(ws, hr, first, last, pcol, vcol, unpriced):
    end_col = max(pcol, vcol or 0, max((c.column for c in ws[hr] if c.value is not None), default=pcol))
    style_table(ws, hr, last, end_col)
    for col in range(1, end_col + 1):
        h = str(ws.cell(hr, col).value or "").lower()
        width = 24 if col in (pcol, vcol) else 12 if h in ("grade", "offer", "qty") else 34 if h in ("cpu", "type") else 44 if "model" in h else 26 if "serial" in h else 19
        ws.column_dimensions[get_column_letter(col)].width = width
    total_col = vcol or pcol
    letter = get_column_letter(total_col)
    # Never overwrite vendor notes, formulas or existing totals below the imported assets.
    last_content = max((c.row for row in ws.iter_rows() for c in row if c.value is not None), default=last)
    start = max(last, last_content) + 2
    for r, label, formula, fill, color in (
        (start, "GRAND TOTAL (INR, ex-GST)", f"=SUM({letter}{first}:{letter}{last})", NAVY, "FFFFFF"),
        (start + 1, "FILTERED TOTAL (INR, ex-GST)", f"=SUBTOTAL(109,{letter}{first}:{letter}{last})", LIME, NAVY),
    ):
        for c in ws.iter_rows(min_row=r, max_row=r, max_col=end_col):
            for cell in c:
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(name="Calibri", size=11, bold=True, color=color)
                cell.alignment = Alignment(vertical="center")
        ws.cell(r, max(1, total_col - 2), label)
        ws.cell(r, total_col, formula).number_format = MONEY
        ws.row_dimensions[r].height = 30
    ws.cell(start + 3, 1, f"Pricing coverage: {last-first+1-unpriced:,} priced rows; {unpriced:,} unpriced. Blank prices are excluded from totals. Historical / provisional rates require review.")
    ws.cell(start + 3, 1).font = Font(name="Calibri", size=10, color="757575", italic=True)
    ws.print_options.horizontalCentered = True
    return [start, start + 1, start + 3], [(f"{letter}{start}"), (f"{letter}{start+1}")]
