"""Excel and CSV exports. Text cells are neutralised against formula injection."""
from __future__ import annotations

import csv
import io

from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="291F56")  # Casterly indigo
LOGO = Path(__file__).resolve().parents[1] / "assets" / "casterly-logo.png"


def add_logo(ws, anchor: str, height_px: int = 96) -> None:
    """Place the Casterly logo on a sheet. Used on report and notes sheets only, never over data tables."""
    if not LOGO.exists():
        return
    img = XLImage(str(LOGO))
    ratio = img.width / img.height
    img.height = height_px
    img.width = round(height_px * ratio)
    ws.add_image(img, anchor)
HEADER_FONT = Font(color="FFFFFF", bold=True)


def safe(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    if isinstance(v, (list, dict)):
        return safe(str(v))
    return v


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for r in rows:
        ws.append([safe(v) for v in r])
    for i, h in enumerate(header, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    ws.freeze_panes = "A2"
    return ws


def to_bytes(wb: Workbook) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def table_csv(header: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    for r in rows:
        w.writerow([safe(v) for v in r])
    return buf.getvalue().encode("utf-8-sig")


SUMMARY_COLS = [
    ("Valuation ID", "id"), ("Item ID", "inventory_item_id"), ("Item", "label"), ("Item ref", "item_ref"),
    ("Lot", "lot_ref"), ("Synthetic demo", "is_synthetic"), ("Valued at", "created_at"), ("As of", "as_of"),
    ("Status", "status"), ("Evidence rating (not a probability)", "rating"), ("Channel", "channel"),
    ("Selling basis", "selling_basis"), ("Sale P25 (incl GST)", "selling_p25"), ("Sale median (incl GST)", "selling_p50"),
    ("Sale P75 (incl GST)", "selling_p75"), ("Recommended listing (incl GST)", "listing_price"),
    ("Target buy / unit", "target_buy"), ("Max safe buy / unit", "max_safe_buy"),
    ("Recommended buy / unit", "recommended_buy"), ("Expected margin", "expected_margin"),
    ("Expected contribution / unit", "expected_contribution_unit"), ("Quantity", "quantity"),
    ("Recommended buy total", "recommended_buy_total"), ("Max safe buy total", "max_safe_buy_total"),
    ("Expected contribution total", "expected_contribution_total"), ("Total lot value (ex-GST revenue)", "total_lot_value"),
    ("Asking total", "asking_total"), ("Asking verdict", "asking_verdict"), ("Flags", "flags"),
    ("Comparables included", "comparables_included"), ("Comparables considered", "comparables_considered"),
    ("Output hash", "output_hash"),
]


def summary_rows(summaries: list[dict]) -> tuple[list[str], list[list]]:
    header = [c[0] for c in SUMMARY_COLS]
    rows = [[(", ".join(s[k]) if isinstance(s.get(k), list) else s.get(k)) for _, k in SUMMARY_COLS] for s in summaries]
    return header, rows


def valuations_workbook(summaries: list[dict]) -> bytes:
    wb = Workbook()
    h, rows = summary_rows(summaries)
    _sheet(wb, "Valuations", h, rows, first=True)
    notes = _sheet(wb, "Notes", ["Note"], [
        ["Prices are INR. Selling prices include GST; buy prices follow the configured tax mode's purchase basis."],
        ["Margin = contribution / sale revenue (ex-GST). Markup = contribution / total cost; they are different."],
        ["The evidence rating is a transparent points score, not a calibrated probability."],
        ["Rows marked 'Synthetic demo' are generated demo data, not market evidence."],
    ])
    add_logo(notes, "C2")
    return to_bytes(wb)


def valuation_detail_workbook(run_id: int, label: str, result: dict, snapshot_meta: dict) -> bytes:
    wb = Workbook()
    sel = result.get("selling") or {}
    econ = result["economics"]
    lot = result.get("lot") or {}
    summary = [
        ["Valuation ID", run_id], ["Item", label], ["Status", result["status"]],
        ["Synthetic demo", result["subject"]["is_synthetic"]], ["As of", result["as_of"]],
        ["Engine version", result["engine_version"]], ["Config version", snapshot_meta.get("config_version")],
        ["Input hash", result["input_hash"]], ["Output hash", result["output_hash"]], ["Channel", result["channel"]],
        ["Selling basis", sel.get("basis")], ["Basis note", sel.get("basis_note")],
        ["Sale P10 (incl GST)", sel.get("p10_gross")], ["Sale P25 (incl GST)", sel.get("p25_gross")],
        ["Sale median (incl GST)", sel.get("p50_gross")], ["Sale P75 (incl GST)", sel.get("p75_gross")],
        ["Sale P90 (incl GST)", sel.get("p90_gross")],
        ["Recommended listing (incl GST)", (result.get("listing_recommendation") or {}).get("price_gross")],
        ["Target buy / unit", econ.get("target_buy")], ["Max safe buy / unit", econ.get("max_safe_buy")],
        ["Recommended buy / unit", econ.get("recommended_buy")],
        ["Quantity", lot.get("quantity")], ["Recommended buy total", lot.get("recommended_buy_total")],
        ["Expected contribution total", lot.get("expected_contribution_total")],
        ["Total lot value (ex-GST revenue)", lot.get("total_lot_value")],
        ["Evidence rating (not a probability)", (result.get("confidence") or {}).get("grade")],
    ]
    summary += [["Message", m] for m in result["messages"]] + [["Needs input", m] for m in result["needs_input"]]
    summary += [["Flag", f["text"]] for f in econ.get("flags", [])]
    ws = _sheet(wb, "Summary", ["Field", "Value"], summary, first=True)
    add_logo(ws, "D2")
    econ_rows = []
    for key, title in (("at_recommended", "At recommended buy, median sale"),
                       ("at_ceiling_conservative", "At max safe buy, P25 sale")):
        if econ.get(key):
            for ln in econ[key]["lines"]:
                econ_rows.append([title, ln["label"], ln["amount"]])
            econ_rows.append([title, "Margin (contribution / revenue)", econ[key]["margin"]])
            econ_rows.append([title, "Markup (contribution / cost)", econ[key]["markup"]])
    _sheet(wb, "Economics", ["Scenario", "Line", "Amount (INR)"], econ_rows)
    inp = econ["inputs"]
    _sheet(wb, "Cost inputs", ["Input", "Value", "Source"], [
        ["Tax mode", inp["tax_mode"], ""], ["GST rate %", inp["gst_rate_pct"], ""],
        ["Target margin %", inp["target_margin_pct"], inp["target_margin_source"]],
        ["Fault category", inp["fault_category"], inp["fault_category_source"]],
        ["Expected yield %", inp["expected_yield_pct"], inp["yield_source"]],
        ["Repair cost / unit", inp["repair_cost_per_unit"], inp["repair_source"]],
        ["Salvage / unsold unit", inp["salvage_per_unsold_unit"], inp["salvage_source"]],
    ] + [[c["name"], c["value"], c["type"] + ("" if c["configured"] else " (NOT CONFIGURED - treated as 0)")]
         for c in inp["cost_lines"]])
    comp_header = ["Evidence ID", "Type", "Channel", "Included", "Exclusion reason", "Set", "Match tier",
                   "Unverified dims", "Adjustments", "Price", "Currency", "Tax basis", "Normalised (incl GST)",
                   "Adjusted price", "Age (days)", "Recency w", "Match w", "Source share", "Final weight",
                   "Seller", "Platform", "Source URL", "Source document", "Source record ID", "Date",
                   "Duplicate of", "Synthetic"]
    comp_rows = []
    for c in result["comparables"]:
        w = c.get("weights") or {}
        comp_rows.append([
            c["id"], c["kind_label"], c["channel"], c["included"], c["exclusion"], c["set"], c["match_tier"],
            ", ".join(c["unverified_dimensions"]),
            "; ".join(f"{a['rule_id']}: {a['dimension_label']} {a['comparable_value']}->{a['subject_value']} "
                      f"{a['method']} {a['value']:+g}{' (UNCALIBRATED)' if not a['calibrated'] else ''}"
                      for a in c["adjustments"]),
            c["price"], c["currency"], c["tax_basis"], c["gross_price"], c["adjusted_price"], c["age_days"],
            w.get("recency"), w.get("match"), w.get("source_share"), w.get("final"), c["seller"], c["platform"],
            c["source_url"], c["source_document"], c["source_record_id"], c["event_date"], c["duplicate_of"],
            c["is_synthetic"]])
    _sheet(wb, "Comparables", comp_header, comp_rows)
    _sheet(wb, "Assumptions", ["Severity", "Code", "Assumption", "Records affected"],
           [[a["severity"], a["code"], a["text"], a.get("count")] for a in result["assumptions"]])
    conf = result.get("confidence")
    if conf:
        _sheet(wb, "Evidence quality", ["Component", "Value", "Points", "Max"],
               [[c["component"], c["value"], c["points"], c["max_points"]] for c in conf["components"]]
               + [["Total", conf["grade"], conf["score"], conf["max_score"]], [conf["note"], "", "", ""]])
    est_rows = [[k, v["count"], v["independent_sources"], v["exact_share"], v["broadened_with_adjustments"],
                 v["p25"], v["p50"], v["p75"], v["dispersion_iqr_over_median"], v["weighted_median_age_days"],
                 v["sufficient"]] for k, v in result["estimates"].items()]
    _sheet(wb, "Evidence sets", ["Set", "Count", "Sources", "Exact share", "Broadened", "P25", "Median", "P75",
                                 "IQR/median", "Median age", "Sufficient"], est_rows)
    return to_bytes(wb)
