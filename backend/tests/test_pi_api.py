"""Pricing Intelligence end-to-end tests: import -> research -> price -> review -> quotation -> export."""
import io
import json
from decimal import Decimal
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook, load_workbook

from .conftest import make_user

FIXTURES = Path(__file__).parent / "fixtures" / "reference"
B002 = FIXTURES / "CIN-20261026-B002_Market_Unit_Prices_Researched.xlsx"
B001_ASSETS = FIXTURES / "CIN-20261026-B001_Asset List.xlsx"
B001_PRICED = FIXTURES / "CIN-20261026-B001_Used_Asset_Market_Prices.xlsx"

B001_HEADER = ["Sales Order", "Global/Local", "Item Code", "Open Description", "Location", "Net Leased Qty",
               "Product Type", "Model", "Processor Type", "Generation", "RAM", "Storage (HDD/SSD)", "Notes"]


def b001_like_workbook() -> bytes:
    """A structural B001-layout fixture (synthetic rows, NOT the real B001 contents)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Asset List"
    ws.append(["Customer asset list — synthetic test fixture"])
    ws.append([])
    ws.append(B001_HEADER)
    rows = [
        ["SO-1", "Local", "IC-01", "Dell Latitude 5430 i7 12th Gen 16GB 512GB", "Pune", 40, "Laptop", "Dell Latitude 5430",
         "Core i7", "12th Gen", "16GB", "512GB SSD", None],
        ["SO-1", "Local", "IC-02", "Samsung Galaxy A54 5G 8GB/128GB", "Pune", 862, "Mobile", "Samsung Galaxy A54 5G 8GB/128GB",
         None, None, None, None, "Bulk handsets"],
        ["SO-2", "Global", "IC-03", "Apple MacBook Pro 16 M1", "Mumbai", 3, "Laptop", "Apple MacBook Pro 16-inch M1",
         "Apple M1", None, "16GB", "512GB SSD", None],
        ["SO-2", "Global", "IC-04", "ThinkPad P1 Gen 4", "Mumbai", "", "Laptop", "Lenovo ThinkPad P1 Gen 4", "Core i7",
         "8th Gen", "32GB", "1TB SSD", None],
        [None, None, None, None, None, 905, None, None, None, None, None, None, None],
        ["Grand Total", None, None, None, None, 905, None, None, None, None, None, None, None],
    ]
    for r in rows:
        ws.append(r)
    notes = wb.create_sheet("Notes")
    notes.append(["Prepared for valuation"])
    notes["B2"] = "=1+1"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def upload(client, h, content, name="vendor.xlsx", **data):
    return client.post("/api/pi/batches", headers=h, files={"file": (name, content)}, data=data)


def test_b001_layout_import_price_override_quote_export(client, admin):
    r = upload(client, admin, b001_like_workbook(), vendor="Acme Leasing", default_grade="B")
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["layout"] == "b001" and b["sheet_name"] == "Asset List" and b["header_row"] == 3
    assert b["lines"] == 4 and b["qty_column"] == "Net Leased Qty" and b["price_column"] is None
    skipped = {s["reason"] for s in b["detection"]["skipped_rows"]}
    assert "total / summary row" in skipped
    assert b["units"] == 40 + 862 + 3  # the blank-quantity line is not counted

    exc = client.get("/api/pi/exceptions", headers=admin, params={"batch_id": b["id"]}).json()
    found = set(exc["summary"])
    assert {"missing_quantity", "cpu_generation_conflict", "apple_chip_unknown", "apple_lock_unknown",
            "large_lot", "exact_cpu_unknown"} <= found

    # No evidence yet: nothing can be priced, nothing is invented
    r = client.post(f"/api/pi/batches/{b['id']}/price", headers=admin)
    assert r.status_code == 200 and set(r.json()["statuses"]) == {"manual_required"}

    groups = client.get(f"/api/pi/batches/{b['id']}/groups", headers=admin).json()
    a54 = next(g for g in groups if "Galaxy A54" in g["label"])
    lat = next(g for g in groups if "Latitude 5430" in g["label"])
    assert a54["units"] == 862

    # A verified supplier quote for the A54, and a refurbished retail listing for the Latitude (not convertible yet)
    r = client.post("/api/pi/observations", headers=admin, json={
        "product_id": a54["product_id"], "price_type": "buyback_quote", "price": "11800", "tax_basis": "inclusive",
        "url": "https://www.cashify.in/sell-old-mobile-phone/used-samsung-galaxy-a54", "seller": "Cashify"})
    assert r.status_code == 200, r.text
    client.post("/api/pi/observations", headers=admin, json={
        "product_id": lat["product_id"], "price_type": "refurb_retail", "price": "41300", "tax_basis": "inclusive"})

    # Configure the A54 500+ band discount and grade B adjustment through a new rules version
    rules = client.get("/api/pi/rules", headers=admin).json()["active"]["rules"]
    rules["categories"]["laptop"]["conversion"]["refurb_retail"] = None
    sp = rules["categories"]["smartphone"]
    sp["grade_adjustment_pct"]["B"] = "10"
    for band in sp["quantity_bands"]:
        if band["min_qty"] == 500:
            band["discount_pct"] = "8"
    r = client.post("/api/pi/rules", headers=admin, json={"rules": rules, "note": "A54 bulk test configuration"})
    assert r.status_code == 200, r.text

    r = client.post(f"/api/pi/batches/{b['id']}/price", headers=admin)
    statuses = r.json()["statuses"]
    assets = client.get(f"/api/pi/batches/{b['id']}/assets", headers=admin).json()["items"]
    a54_asset = next(a for a in assets if a["product_id"] == a54["product_id"])
    # 11800 incl. GST -> 10000 ex-GST, x1.0 (buyback quote) x0.90 (grade B) x0.92 (500+ units) = 8280 -> 8250
    assert a54_asset["calculated_price"] == "8250.00" and a54_asset["total_value"] == str(Decimal("8250") * 862) + ".00"
    lat_asset = next(a for a in assets if a["product_id"] == lat["product_id"])
    assert lat_asset["calculated_price"] is None and lat_asset["status"] == "manual_required"
    assert statuses.get("auto_priced", 0) + statuses.get("review", 0) >= 1

    detail = client.get(f"/api/pi/assets/{a54_asset['id']}", headers=admin).json()
    steps = [s["step"] for s in detail["estimate"]["breakdown"]["steps"]]
    assert steps[0] == "benchmark" and detail["estimate"]["benchmark"] == "10000.00"

    # Overrides need a reason and the sales role cannot override
    assert client.post("/api/pi/assets/override", headers=admin,
                       json={"asset_ids": [lat_asset["id"]], "price": "26000", "reason": ""}).status_code == 422
    sales = make_user(client, admin, "sales")
    assert client.post("/api/pi/assets/override", headers=sales,
                       json={"asset_ids": [lat_asset["id"]], "price": "26000", "reason": "Vendor"}).status_code == 403
    r = client.post("/api/pi/assets/override", headers=admin,
                    json={"asset_ids": [lat_asset["id"]], "price": "26000", "reason": "Matched last month's purchase"})
    assert r.status_code == 200
    d = client.get(f"/api/pi/assets/{lat_asset['id']}", headers=admin).json()
    assert d["asset"]["calculated_price"] is None and d["asset"]["override_price"] == "26000.00"
    assert d["overrides"][0]["reason"] == "Matched last month's purchase"

    # Approve and quote; retrying does not create a duplicate quotation
    assert client.post(f"/api/pi/batches/{b['id']}/quotations", headers=admin, json={}).status_code == 422
    r = client.post("/api/pi/assets/approve", headers=admin, json={"asset_ids": [a54_asset["id"], lat_asset["id"]]})
    assert r.json()["approved"] == 2
    q1 = client.post(f"/api/pi/batches/{b['id']}/quotations", headers=admin, json={"vendor": "Acme"}).json()
    q2 = client.post(f"/api/pi/batches/{b['id']}/quotations", headers=admin, json={"vendor": "Acme"}).json()
    assert q1["id"] == q2["id"] and q1["created"] and not q2["created"]
    assert Decimal(q1["total"]) == Decimal("8250") * 862 + Decimal("26000") * 40

    # Export: original workbook preserved, Market Unit Price + Estimated Lot Value appended, sheets added
    r = client.get(f"/api/pi/batches/{b['id']}/export", headers=admin)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:2] == ["Asset List", "Notes"] and wb["Notes"]["B2"].value == "=1+1"
    for s in ("Valuation Summary", "Model Pricing Evidence", "Review Before Purchase", "Rate Benchmarks", "Research Sources"):
        assert s in wb.sheetnames
    ws = wb["Asset List"]
    hdr = [c.value for c in ws[3]]
    assert hdr[:13] == B001_HEADER and hdr[13] == "Market Unit Price" and hdr[14] == "Estimated Lot Value"
    a54_row = a54_asset["row_number"]
    assert ws.cell(row=a54_row, column=14).value == 8250
    assert ws.cell(row=a54_row, column=15).value.startswith("=IF(ISNUMBER(N")
    assert ws["A1"].value.startswith("Customer asset list")
    q = client.get(f"/api/pi/quotations/{q1['id']}/export", headers=admin)
    assert q.status_code == 200 and "Vendor Quotation" in load_workbook(io.BytesIO(q.content)).sheetnames


def test_manual_valuation_and_mapping(client, admin):
    r = client.post("/api/pi/batches/manual", headers=admin, json={"name": "Quick check", "assets": [
        {"description": "Intel Xeon Gold 6240", "model": "XEON GOLD 6240", "category": "server_cpu", "grade": "B",
         "quantity": 4}]})
    assert r.status_code == 200, r.text
    assert r.json()["layout"] == "manual_entry" and r.json()["units"] == 4

    wb = Workbook()
    ws = wb.active
    ws.append(["Thing", "How many", "Kit"])
    ws.append(["Dell Latitude 7440 i5-1345U 16GB 512GB", 3, "Laptop"])
    buf = io.BytesIO()
    wb.save(buf)
    r = upload(client, admin, buf.getvalue(), name="odd.xlsx")
    assert r.status_code == 422  # unrecognised layout: manual mapping required
    # map manually: first create the batch with an explicit sheet/header, then remap
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(["Description", "Qty"])
    ws2.append(["Dell Latitude 7440 i5-1345U 16GB 512GB", 3])
    buf2 = io.BytesIO()
    wb2.save(buf2)
    b = upload(client, admin, buf2.getvalue(), name="simple.xlsx").json()
    assert b["lines"] == 1 and b["units"] == 3
    r = client.post(f"/api/pi/batches/{b['id']}/mapping", headers=admin,
                    json={"mapping": {"description": "Description", "quantity": "Qty"}})
    assert r.status_code == 200 and r.json()["lines"] == 1


class FakeSearch:
    """httpx transport standing in for the Brave Search API in tests (not real market data)."""

    def __init__(self, status=200, results=None):
        self.status, self.results, self.calls = status, results or [], 0

    def handler(self, request: httpx.Request) -> httpx.Response:
        self.calls += 1
        if self.status != 200:
            return httpx.Response(self.status, json={})
        return httpx.Response(200, json={"web": {"results": self.results}})


@pytest.fixture()
def brave_key(monkeypatch):
    from app.config import get_settings
    monkeypatch.setenv("CLPI_BRAVE_API_KEY", "test-key")
    get_settings.cache_clear()
    yield
    monkeypatch.delenv("CLPI_BRAVE_API_KEY")
    get_settings.cache_clear()


def _cpu_batch(client, admin):
    r = client.post("/api/pi/batches/manual", headers=admin, json={"name": "CPU lot", "assets": [
        {"description": "Intel Xeon Gold 6130", "model": "XEON GOLD 6130", "category": "server_cpu", "grade": "B",
         "quantity": 20}]})
    return r.json()


def test_research_job_candidates_verification_and_cache(client, admin, brave_key):
    from app.procurement import research
    from app.procurement.rules import default_rules
    b = _cpu_batch(client, admin)
    assert client.post(f"/api/pi/batches/{b['id']}/research", headers=admin, json={"provider": "tavily"}).status_code == 409
    job = client.post(f"/api/pi/batches/{b['id']}/research", headers=admin, json={}).json()
    assert job["provider"] == "brave" and job["total"] == 1
    again = client.post(f"/api/pi/batches/{b['id']}/research", headers=admin, json={}).json()
    assert again["id"] == job["id"]  # one live job per batch
    fake = FakeSearch(results=[
        {"title": "Intel Xeon Gold 6130 2.1GHz 16-Core Processor (Refurbished)", "url": "https://servercart.in/p/6130",
         "description": "Refurbished Intel Xeon Gold 6130 SR3B9 ₹6,900 incl. GST, 6 months warranty"},
        {"title": "Intel Xeon Gold 6148 Processor", "url": "https://servercart.in/p/6148",
         "description": "Xeon Gold 6148 ₹4,869"},
        {"title": "Xeon Gold 6130 used", "url": "https://www.olx.in/item/6130", "description": "Used server CPU, call for price"},
    ])
    rules = default_rules()
    rules.research.concurrency = 1
    research.run_job(job["id"], rules, transport=httpx.MockTransport(fake.handler))
    j = client.get(f"/api/pi/jobs/{job['id']}", headers=admin).json()
    assert j["status"] == "completed" and j["observations_found"] == 1  # same listing in both queries is deduplicated
    assert j["tasks"][0]["queries"][1]["rejected"]["duplicate listing"] == 3  # rejected records are also retained/deduplicated
    assert j["tasks"][0]["queries"][0]["rejected"] == {"not the same product (technical identity mismatch)": 1}
    obs = client.get("/api/pi/observations", headers=admin, params={"batch_id": b["id"]}).json()
    priced = [o for o in obs if o["status"] == "candidate"]
    leads = [o for o in obs if o["status"] == "lead"]
    assert {o["price"] for o in priced} == {"6900.00"} and leads and all(o["price"] is None for o in leads)
    assert priced[0]["price_type"] == "refurb_retail" and priced[0]["url"] == "https://servercart.in/p/6130"

    # Candidates are not used until verified
    client.post(f"/api/pi/batches/{b['id']}/price", headers=admin)
    a = client.get(f"/api/pi/batches/{b['id']}/assets", headers=admin).json()["items"][0]
    assert a["status"] == "manual_required"
    client.post("/api/pi/observations/verify", headers=admin, json={"ids": [priced[0]["id"]], "action": "verify",
                                                                    "tax_basis": "inclusive"})
    rules_doc = client.get("/api/pi/rules", headers=admin).json()["active"]["rules"]
    rules_doc["categories"]["server_cpu"]["conversion"]["refurb_retail"] = "0.62"
    rules_doc["categories"]["server_cpu"]["grade_adjustment_pct"]["B"] = "0"
    rules_doc["categories"]["server_cpu_xeon"]["conversion"]["refurb_retail"] = "0.62"
    rules_doc["categories"]["server_cpu_xeon"]["grade_adjustment_pct"]["B"] = "0"
    for band in rules_doc["categories"]["server_cpu_xeon"]["quantity_bands"]:
        band["discount_pct"] = "0"
    client.post("/api/pi/rules", headers=admin, json={"rules": rules_doc, "note": "CPU retail conversion"})
    client.post(f"/api/pi/batches/{b['id']}/price", headers=admin)
    a = client.get(f"/api/pi/batches/{b['id']}/assets", headers=admin).json()["items"][0]
    # 6900 / 1.18 = 5847.46 x 0.62 = 3625.42 -> round down to 50 = 3600
    assert a["calculated_price"] == "3600.00" and a["benchmark"] == "3625.42"

    # Re-running research inside the cache window does not search again or duplicate evidence
    calls = fake.calls
    job2 = client.post(f"/api/pi/batches/{b['id']}/research", headers=admin, json={}).json()
    research.run_job(job2["id"], rules, transport=httpx.MockTransport(fake.handler))
    assert fake.calls == calls
    assert client.get(f"/api/pi/jobs/{job2['id']}", headers=admin).json()["cached"] == 1


def test_research_failure_retry_and_cancel(client, admin, brave_key):
    from app.procurement import research
    from app.procurement.rules import default_rules
    b = _cpu_batch(client, admin)
    rules = default_rules()
    rules.research.retry_backoff_seconds = 0
    rules.research.max_attempts = 2
    rules.research.concurrency = 1
    fake = FakeSearch(status=503)
    job = client.post(f"/api/pi/batches/{b['id']}/research", headers=admin, json={}).json()
    research.run_job(job["id"], rules, transport=httpx.MockTransport(fake.handler), wait_for_backoff=True)
    j = client.get(f"/api/pi/jobs/{job['id']}", headers=admin).json()
    assert j["status"] == "failed" and j["tasks"][0]["attempts"] == 2 and "503" in j["tasks"][0]["error"]
    # resume after the provider recovers
    assert client.post(f"/api/pi/jobs/{job['id']}/resume", headers=admin).json()["status"] == "queued"
    fake.status = 200
    research.run_job(job["id"], rules, transport=httpx.MockTransport(fake.handler))
    assert client.get(f"/api/pi/jobs/{job['id']}", headers=admin).json()["status"] == "completed"
    # cancel a queued job
    job3 = client.post(f"/api/pi/batches/{b['id']}/research", headers=admin, json={"force": True}).json()
    assert client.post(f"/api/pi/jobs/{job3['id']}/cancel", headers=admin).json()["status"] == "cancelled"
    # 401 is not retried
    fake.status = 401
    job4 = client.post(f"/api/pi/batches/{b['id']}/research", headers=admin, json={"force": True}).json()
    research.run_job(job4["id"], rules, transport=httpx.MockTransport(fake.handler), wait_for_backoff=True)
    j4 = client.get(f"/api/pi/jobs/{job4['id']}", headers=admin).json()
    assert j4["status"] == "failed" and j4["tasks"][0]["attempts"] == 1


def test_no_provider_configured_is_reported(client, admin):
    b = _cpu_batch(client, admin)
    r = client.post(f"/api/pi/batches/{b['id']}/research", headers=admin, json={})
    assert r.status_code == 409 and "CLPI_BRAVE_API_KEY" in r.json()["detail"]
    meta = client.get("/api/pi/meta", headers=admin).json()
    assert meta["default_provider"] is None and not any(p["configured"] for p in meta["providers"])


def test_outcomes_and_accuracy(client, admin):
    b = _cpu_batch(client, admin)
    pid = client.get(f"/api/pi/batches/{b['id']}/groups", headers=admin).json()[0]["product_id"]
    client.post("/api/pi/observations", headers=admin, json={"product_id": pid, "price_type": "supplier_quote",
                                                             "price": "4300", "tax_basis": "exclusive"})
    client.post(f"/api/pi/batches/{b['id']}/price", headers=admin)
    from datetime import date
    r = client.post("/api/pi/outcomes", headers=admin, json={"product_id": pid, "kind": "purchase", "unit_price": "4000",
                                                             "quantity": 20, "outcome_date": date.today().isoformat()})
    assert r.status_code == 200
    rep = client.get("/api/pi/accuracy", headers=admin).json()
    assert rep["purchase"]["overall"]["n"] == 1
    assert rep["purchase"]["overall"]["mae"] == 250.0 and rep["purchase"]["overall"]["bias_pct"] == -6.25  # 4300 x .9 x .98 -> 3750 starter buy price
    hist = client.get(f"/api/pi/products/{pid}", headers=admin).json()
    assert hist["outcomes"] and any(o["price_type"] == "verified_transaction" for o in hist["observations"])


def test_sources_and_rules_permissions(client, admin):
    srcs = client.get("/api/pi/sources", headers=admin).json()
    assert {"servercart.in", "cashify.in", "serverbasket.com"} <= {s["domain"] for s in srcs}
    proc = make_user(client, admin, "procurement")
    assert client.post("/api/pi/rules", headers=proc, json={"rules": {}, "note": "nope"}).status_code == 403
    assert client.post("/api/pi/sources", headers=proc, json={"name": "X", "domain": "x.in"}).status_code == 403
    r = client.post("/api/pi/rules/presets/historical_reference", headers=admin)
    assert r.status_code == 200 and not r.json()["is_active"]
    bad = client.post("/api/pi/rules", headers=admin, json={"rules": {"gst_pct": "abc"}, "note": "bad rules"})
    assert bad.status_code == 422


# --------------------------------------------------------------------------- B002 regression (real workbook)

@pytest.mark.skipif(not B002.exists(), reason="B002 reference workbook not present in tests/fixtures/reference")
def test_b002_regression(client, admin):
    content = B002.read_bytes()
    original = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    hist = {}
    for i, row in enumerate(original["Sheet1"].iter_rows(values_only=True), start=1):
        if i > 1 and row[0] is not None:
            hist[i] = (row[0], row[1], row[9])
    original.close()
    assert len(hist) == 2892
    assert sum(1 for v in hist.values() if v[1] == "RC-RAM") == 2698
    assert sum(1 for v in hist.values() if v[1] == "PM-CPU") == 194

    b = upload(client, admin, content, name="CIN-20261026-B002_Asset List.xlsx").json()
    assert b["layout"] == "b002" and b["lines"] == 2892 and b["units"] == 2892 and b["price_column"] == "Market Unit Price"
    groups = client.get(f"/api/pi/batches/{b['id']}/groups", headers=admin).json()
    assert len(groups) == 52  # 51 identified product groups + 1 unidentified (XEON SILVER 41112)
    assert all(g["grade"] == "B" for g in groups)
    gold = next(g for g in groups if g["label"] == "Intel Xeon Gold 6130")
    assert gold["units"] == 20
    exc = client.get("/api/pi/exceptions", headers=admin, params={"batch_id": b["id"]}).json()["summary"]
    assert exc["ram_speed_conflict"] == 235 and exc["possible_typo"] == 1 and exc["cpu_sku_inferred"] == 7

    # Historical configuration: load the dated reference valuation and the regression rule preset
    r = client.post("/api/pi/reference-packs", headers=admin, files={"file": (B002.name, content)})
    assert r.status_code == 200, r.text
    ref = r.json()
    assert ref["reference_date"] == "2026-10-26" and ref["historical_observations"] == 51 and ref["price_conflicts"] == 0
    preset = client.post("/api/pi/rules/presets/historical_reference", headers=admin).json()
    r = client.post(f"/api/pi/batches/{b['id']}/price", headers=admin, params={"rule_version_id": preset["id"]})
    assert r.json()["statuses"] == {"provisional": 2891, "manual_required": 1}

    out = client.get(f"/api/pi/batches/{b['id']}/export", headers=admin)
    assert out.headers["x-export-prices-written"] == "2891"
    wb = load_workbook(io.BytesIO(out.content), read_only=True)
    ws = wb["Sheet1"]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    assert header[9] == "Market Unit Price" and len([h for h in header if h]) == 10  # column J filled, none added
    mismatches, ram_total, cpu_total = [], Decimal(0), Decimal(0)
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i not in hist:
            continue
        serial, cat, old = hist[i]
        assert row[0] == serial  # serial-to-row mapping preserved
        new = row[9]
        if new is None:
            mismatches.append((i, serial))
            continue
        if Decimal(str(new)) != Decimal(str(old)):
            mismatches.append((i, serial, old, new))
        if cat == "RC-RAM":
            ram_total += Decimal(str(new))
        else:
            cpu_total += Decimal(str(new))
    wb.close()
    assert mismatches == [(49, "NA-744")]  # 'XEON SILVER 41112' is held for manual pricing, not guessed as 4112
    assert ram_total == Decimal("12142800") and cpu_total == Decimal("1191300")
    assert ram_total + cpu_total + Decimal("3000") == Decimal("13337100")  # historical total incl. the held row
    names = set(load_workbook(io.BytesIO(out.content), read_only=True).sheetnames)
    # A user confirms the typo with the vendor and corrects the identification; the row joins the 4112 group
    typo = client.get(f"/api/pi/batches/{b['id']}/assets", headers=admin, params={"q": "41112"}).json()["items"][0]
    r = client.post(f"/api/pi/assets/{typo['id']}/reidentify", headers=admin,
                    json={"fields": {"model": "XEON SILVER 4112"}, "method": "manual"})
    assert r.status_code == 200 and r.json()["asset"]["label"] == "Intel Xeon Silver 4112"
    client.post(f"/api/pi/batches/{b['id']}/price", headers=admin, params={"rule_version_id": preset["id"]})
    assert Decimal(client.get(f"/api/pi/batches/{b['id']}", headers=admin).json()["estimated_value"]) == Decimal("13337100")
    assert {"Valuation Summary", "Model Pricing Evidence", "Review Before Purchase", "Rate Benchmarks (CLPI)",
            "Research Sources (CLPI)", "Pricing Summary", "Rate Benchmarks", "Model Valuation"} <= names


@pytest.mark.skipif(not (B001_ASSETS.exists() and B001_PRICED.exists()),
                    reason="B001 reference workbooks not present in tests/fixtures/reference")
def test_b001_regression(client, admin):  # pragma: no cover - runs only where the confidential files exist
    b = upload(client, admin, B001_ASSETS.read_bytes(), name=B001_ASSETS.name).json()
    assert b["layout"] == "b001" and b["lines"] == 42 and b["units"] == 1324
    groups = client.get(f"/api/pi/batches/{b['id']}/groups", headers=admin).json()
    assert any(g["units"] == 862 and "A54" in g["label"] for g in groups)
    r = client.post("/api/pi/reference-packs", headers=admin, files={"file": (B001_PRICED.name, B001_PRICED.read_bytes())},
                    data={"reference_date": "2026-10-26"})
    assert r.status_code == 200, r.text
    preset = client.post("/api/pi/rules/presets/historical_reference", headers=admin).json()
    client.post(f"/api/pi/batches/{b['id']}/price", headers=admin, params={"rule_version_id": preset["id"]})
    total = Decimal(client.get(f"/api/pi/batches/{b['id']}", headers=admin).json()["estimated_value"])
    assert total == Decimal("20260000"), json.dumps(str(total))
