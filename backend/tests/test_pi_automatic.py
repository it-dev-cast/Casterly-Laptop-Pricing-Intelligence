"""Real API/worker/export integration; external responses explicitly simulated, never a live-search claim."""
import io
import hashlib
import json
import shutil
import pytest
from datetime import datetime, timezone
from decimal import Decimal
from zipfile import ZipFile
import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from PIL import Image as PILImage
from app.config import get_settings
from app.db import SessionLocal
from app.procurement import automatic, worker
from app.procurement.providers import SearchHit, BraveProvider, TavilyProvider, ProviderError
from app.procurement.extraction import confirmed_offer, to_candidate, match_quality
from app.procurement.identify import identify
from app.procurement.rules import default_rules, legacy_default_rules
from app.procurement.engine import Group, Obs, price_group
from .test_pi_api import upload, b001_like_workbook, B002


def sheet_bytes(rows, name='Asset List'):
    wb=Workbook();ws=wb.active;ws.title=name
    for row in rows: ws.append(row)
    out=io.BytesIO();wb.save(out);return out.getvalue()


def finish(client,admin,batch):
    r=client.post(f"/api/pi/batches/{batch['id']}/analyze",headers=admin)
    assert r.status_code==200,r.text
    j=r.json()
    assert client.post(f"/api/pi/batches/{batch['id']}/analyze",headers=admin).json()['id']==j['id']
    assert worker.run_once()==j['id']
    state=client.get(f"/api/pi/jobs/{j['id']}",headers=admin).json()
    assert state['status'] in ('completed','completed_with_errors'),state
    artifact=client.get(f"/api/pi/jobs/{j['id']}/workbook",headers=admin)
    assert artifact.status_code==200,artifact.text
    return state,artifact.content


def test_starter_policy_is_installed_on_clean_database(client,admin):
    doc=client.get('/api/pi/rules',headers=admin).json()
    rules=doc['active']['rules']
    assert 'starter' in rules['name'].lower() and not doc['gaps']
    for cat in ('laptop','desktop','monitor','macbook','smartphone','server_ram_ddr3','server_ram_ddr4','server_ram_ddr5','server_cpu_xeon','server_cpu_epyc','mobile_workstation','other'):
        assert rules['categories'][cat]['conversion']['refurb_retail']
        assert rules['categories'][cat]['grade_adjustment_pct']['Untested']=='35'
    assert 'assumptions' in rules['policy_basis']


def test_offline_one_click_no_fabricated_prices(client,admin):
    original=b001_like_workbook()
    b=upload(client,admin,original).json()
    j,content=finish(client,admin,b)
    assert j['workflow']['result']=='Research Unavailable'
    assert j['workflow']['research_mode']=='offline_history'
    assert j['workflow']['priced_lines']==0
    assert [e['phase'] for e in j['workflow']['events']]==automatic.PHASES+['Research Unavailable']
    assets=client.get(f"/api/pi/batches/{b['id']}/assets",headers=admin).json()['items']
    assert all(a['final_price'] is None for a in assets)
    wb=load_workbook(io.BytesIO(content));assert wb['Asset List']['N3'].value=='Market Unit Price'
    assert wb['Notes']['B2'].value=='=1+1'
    with SessionLocal() as db:
        from app.procurement.models import PIBatch
        from app.importing.files import stored_path
        assert hashlib.sha256(stored_path(db.get(PIBatch,b['id']).stored_name).read_bytes()).digest()==hashlib.sha256(original).digest()


def test_configured_history_automatic_reuse_b002(client,admin,monkeypatch,tmp_path):
    assert B002.exists()
    history=tmp_path/'completed-history';history.mkdir();shutil.copy2(B002,history/B002.name)
    monkeypatch.setenv('CLPI_REFERENCE_DIRS',json.dumps([str(history)]));get_settings.cache_clear()
    source=load_workbook(B002)
    ws=source['Sheet1']
    historical_rows={r:(ws.cell(r,1).value,ws.cell(r,2).value,ws.cell(r,10).value) for r in range(2,ws.max_row+1) if ws.cell(r,1).value is not None}
    assert sum(v[1]=='RC-RAM' for v in historical_rows.values())==2698
    assert sum(v[1]=='PM-CPU' for v in historical_rows.values())==194
    for row in range(2,ws.max_row+1): ws.cell(row,10).value=None
    out=io.BytesIO();source.save(out)
    b=upload(client,admin,out.getvalue(),name='B002-unpriced-derived.xlsx').json()
    assert b['lines']==2892 and b['units']==2892 and b['price_column']=='Market Unit Price'
    j,data=finish(client,admin,b)
    assert j['workflow']['priced_lines']==2891 and j['workflow']['result']=='Needs Review'
    assert j['workflow']['integrity']['prices_written']==2891
    assert j['total']<100
    assert client.get(f"/api/pi/batches/{b['id']}",headers=admin).json()['estimated_value']=='13334100.00'
    assets=client.get(f"/api/pi/batches/{b['id']}/assets?page_size=500",headers=admin).json()['items']
    assert all((a['status']=='provisional' and a['research_methods']==['historical']) or a['status']=='manual_required' for a in assets)
    wb=load_workbook(io.BytesIO(data)); headers=[c.value for c in wb['Sheet1'][1]]
    assert headers[9]=='Market Unit Price' and headers.count('Market Unit Price')==1
    for row,(serial,category,old_price) in historical_rows.items():
        assert wb['Sheet1'].cell(row,1).value==serial
        assert wb['Sheet1'].cell(row,10).value==(None if row==49 else old_price)
    with SessionLocal() as db:
        assert automatic.import_configured_references(db)==[]
    get_settings.cache_clear()


def test_product_matching_rejects_speed_cpu_and_module_conflicts():
    assert match_quality('server_ram',{'capacity_gb':16,'ddr':'DDR4','speed_mts':2400,'module_type':'RDIMM'},'16GB DDR4 2666 ECC RDIMM') is None
    assert match_quality('server_cpu',{'model':'Xeon Gold 6130'},'Intel Xeon Gold 6130P used') is None
    assert match_quality('server_ram',{'capacity_gb':32,'ddr':'DDR4','speed_mts':2666,'module_type':'RDIMM'},'32GB DDR4 2666 ECC LRDIMM') is None
    a=identify({'model':'Dell Latitude 5430','cpu':'i5-1245U','ram':'16GB','storage':'256GB','grade':'A','condition':'tested'})
    b=identify({'model':'Dell Latitude 5430','cpu':'i5-1245U','ram':'16GB','storage':'256GB','grade':'A','condition':'untested'})
    assert a.product_key!=b.product_key and 'untested' in b.risk_flags


def test_snippets_never_confirm_and_ambiguous_prices_reject():
    spec={'model':'Xeon Gold 6130','vendor':'Intel'}
    hit=SearchHit('Used Intel Xeon Gold 6130','https://example.com/item','Price INR 5000','brave')
    assert confirmed_offer(hit,'server_cpu',spec)==(None,None)
    assert to_candidate(hit,'server_cpu',spec).price==5000
    hit.snippet='Lot of 10 used Intel Xeon Gold 6130 price INR 5000'
    assert 'unit amount not confirmed' in to_candidate(hit,'server_cpu',spec).reason
    hit.snippet='Price INR 5000; EMI INR 500'
    assert to_candidate(hit,'server_cpu',spec).reason
    hit.extra={'raw_content':'Used Intel Xeon Gold 6130 price INR 5000'}
    assert confirmed_offer(hit,'server_cpu',spec)==(5000,'licensed_full_content_single_price')
    hit.extra={'raw_content':'Used Intel Xeon Gold 6130 EMI per month INR 5000'}
    assert confirmed_offer(hit,'server_cpu',spec)==(None,None)


def test_simulated_provider_fallback_confirmed_offer(client,admin,monkeypatch):
    monkeypatch.setenv('CLPI_BRAVE_API_KEY','test-brave-key');monkeypatch.setenv('CLPI_TAVILY_API_KEY','test-tavily-key');get_settings.cache_clear()
    calls=[]
    def handler(req):
        calls.append(req.url.host)
        if 'brave' in req.url.host: return httpx.Response(401)
        return httpx.Response(200,json={'results':[{'title':'Used Intel Xeon Gold 6130', 'url':'https://example.com/6130',
            'content':'Intel Xeon Gold 6130 used price INR 5000',
            'raw_content':'Intel Xeon Gold 6130 used price INR 5000'}]})
    from app.procurement.providers import Provider
    monkeypatch.setattr(Provider,'_client',lambda self:httpx.Client(transport=httpx.MockTransport(handler),timeout=2))
    b=upload(client,admin,sheet_bytes([['Model','Category','Grade','Qty'],['XEON GOLD 6130','PM-CPU','A',10]])).json()
    j,data=finish(client,admin,b)
    assert any('brave' in s for s in calls) and any('tavily' in s for s in calls)
    assert j['workflow']['priced_lines']==1
    obs=client.get('/api/pi/observations',headers=admin).json()
    assert obs[0]['provider']=='tavily' and obs[0]['status']=='verified'
    assert obs[0]['attributes']['source_timestamp'] is None
    a=client.get(f"/api/pi/batches/{b['id']}/assets",headers=admin).json()['items'][0]
    assert a['status']=='provisional' and a['research_methods']==['live_search_confirmed']
    get_settings.cache_clear()


def test_complex_workbook_original_parts_retained(client,admin):
    wb=Workbook();ws=wb.active;ws.title='Asset List';ws.append(['Model','Grade','Qty']);ws.append(['Dell Latitude 5430','A',2]);ws['F4']='=1+1'
    chart=BarChart();chart.add_data(Reference(ws,min_col=3,min_row=1,max_row=2),titles_from_data=True);ws.add_chart(chart,'H5')
    pic=io.BytesIO();PILImage.new('RGB',(10,10),'red').save(pic,format='PNG');pic.seek(0);ws.add_image(Image(pic),'H20')
    wb.create_sheet('Valuation Summary')['A1']='Existing customer sheet'
    wb.create_sheet('Valuation Summary (CLPI)')['A1']='Also preserve this'
    stream=io.BytesIO();wb.save(stream);original=stream.getvalue()
    b=upload(client,admin,original).json();j,data=finish(client,admin,b)
    aa,bb=ZipFile(io.BytesIO(original)),ZipFile(io.BytesIO(data))
    modified=set(j['workflow']['integrity']['fidelity']['modified_parts'])
    assert set(aa.namelist())<=set(bb.namelist())
    for name in aa.namelist():
        if name not in modified: assert aa.read(name)==bb.read(name),name
    assert any(n.startswith('xl/charts/') for n in aa.namelist())
    result=load_workbook(io.BytesIO(data))
    assert result['Asset List']['F4'].value=='=1+1'
    assert result['Valuation Summary']['A1'].value=='Existing customer sheet'
    assert result['Valuation Summary (CLPI)']['A1'].value=='Also preserve this'
    assert 'Valuation Summary (CLPI 2)' in result.sheetnames


def test_condition_discount_not_applied_twice():
    rules=default_rules();when=datetime.now(timezone.utc)
    ob=Obs(1,Decimal('10000'),'INR','exclusive','supplier_quote','identical',when,'test',1,'verified')
    r=price_group(Group('laptop','Untested',1,risk_flags=['untested']),[ob],rules,when)
    assert r['recommended_unit']=='6500.00'
    historical=Obs(2,Decimal('10000'),'INR','exclusive','historical_valuation','identical',when,'history',1,'verified')
    r=price_group(Group('laptop','B',100,risk_flags=['battery_low','untested']),[historical],rules,when)
    assert r['recommended_unit']=='10000.00' and r['status']=='provisional'


def test_provider_malformed_response_and_timeout(monkeypatch):
    for cls in (BraveProvider,TavilyProvider):
        try:
            cls(httpx.MockTransport(lambda req:httpx.Response(200,json=[]))).search('test',count=1,country='IN')
            assert False,'Malformed response silently accepted'
        except ProviderError as e: assert e.retryable
        def timeout(req): raise httpx.ReadTimeout('test')
        try:
            cls(httpx.MockTransport(timeout)).search('test',count=1,country='IN')
            assert False
        except ProviderError as e: assert e.retryable and 'key' not in str(e).lower()

def test_approved_catalog_offline_workflow(client,admin,monkeypatch,tmp_path):
    data=sheet_bytes([['Model','Category','Grade','Qty'],['XEON GOLD 6130','PM-CPU','A',1]])
    b=upload(client,admin,data).json()
    catalog=tmp_path/'approved.json'
    catalog.write_text(json.dumps([{'product_key':'server_cpu|xeon-gold-6130','price':'11800','currency':'INR',
        'tax_basis':'inclusive','price_type':'refurb_retail','observed_at':datetime.now(timezone.utc).isoformat(),
        'url':'https://example.com/approved-6130','seller':'Synthetic approved vendor','condition':'used','grade':'A'}]))
    monkeypatch.setenv('CLPI_APPROVED_CATALOG_PATH',str(catalog));monkeypatch.setenv('CLPI_APPROVED_CATALOG_TERMS_CONFIRMED','true');get_settings.cache_clear()
    j,out=finish(client,admin,b)
    a=client.get(f"/api/pi/batches/{b['id']}/assets",headers=admin).json()['items'][0]
    assert a['calculated_price']=='6000.00' and a['research_methods']==['approved_catalog']
    assert j['workflow']['catalog']['added']==1
    clean=client.get(f"/api/pi/batches/{b['id']}/export?clean=true",headers=admin)
    assert clean.status_code==200
    assert 'Clean reserialized' in clean.headers['x-workbook-fidelity']
    get_settings.cache_clear()


def test_diagnostics_worker_liveness_and_blocked_probe(client,admin):
    from app.procurement.diagnostics import heartbeat
    heartbeat()
    d=client.get('/api/pi/diagnostics',headers=admin).json()
    assert d['queue']['worker_healthy'] and d['research_mode']=='offline_history'
    for p in d['providers']: assert p['connectivity_verified'] is None
    result=client.post('/api/pi/providers/brave/probe',headers=admin).json()
    assert result=={'status':'blocked','requirement':'CLPI_BRAVE_API_KEY'}


def test_old_policy_upgrade_preserves_admin_versions(client,admin):
    from app.procurement.service import active_rules, save_rules
    with SessionLocal() as db:
        before=save_rules(db,legacy_default_rules(),None,'Initial defaults: identity conversions only; no discounts configured')
        # Simulate an older JSON document before new optional fields were introduced.
        original=dict(before.rules);original.pop('policy_basis');original.pop('similar_model_factors')
        original['research'].pop('auto_accept_confirmed')
        before.rules=original
        db.commit()
        rules,after=active_rules(db)
        assert after.id!=before.id and rules.policy_basis
        custom=save_rules(db,legacy_default_rules(),None,'Approved administrator custom policy')
        db.commit()
        assert active_rules(db)[1].id==custom.id


def test_catalog_requires_authorization_and_source_date(client,admin,monkeypatch,tmp_path):
    data=sheet_bytes([['Model','Category','Grade','Qty'],['XEON GOLD 6130','PM-CPU','A',1]])
    b=upload(client,admin,data).json()
    catalog=tmp_path/'bad.json';catalog.write_text('[{"product_key":"server_cpu|xeon-gold-6130","price":"5000"}]')
    monkeypatch.setenv('CLPI_APPROVED_CATALOG_PATH',str(catalog));get_settings.cache_clear()
    j,_=finish(client,admin,b)
    assert j['workflow']['catalog']['status']=='not_configured' and j['workflow']['priced_lines']==0
    monkeypatch.setenv('CLPI_APPROVED_CATALOG_TERMS_CONFIRMED','true');get_settings.cache_clear()
    j,_=finish(client,admin,b)
    assert j['workflow']['catalog']['rejected']==1 and j['workflow']['priced_lines']==0
    get_settings.cache_clear()


def test_ooxml_opaque_parts_preserved(client,admin):
    # Synthetic opaque cache/extension members prove ZIP fidelity, not Excel pivot execution.
    source=sheet_bytes([['Model','Grade','Qty'],['Dell Latitude 5430','A',1]])
    stream=io.BytesIO()
    from zipfile import ZIP_DEFLATED
    with ZipFile(io.BytesIO(source)) as old,ZipFile(stream,'w',ZIP_DEFLATED) as new:
        for name in old.namelist():new.writestr(name,old.read(name))
        new.writestr('xl/pivotCache/syntheticOpaqueCache.xml',b'<opaque>Preserve cache bytes</opaque>')
        new.writestr('customXml/syntheticExtension.xml',b'<vendor-extension>Keep untouched</vendor-extension>')
    b=upload(client,admin,stream.getvalue()).json();_,out=finish(client,admin,b)
    with ZipFile(io.BytesIO(out)) as z:
        assert z.read('xl/pivotCache/syntheticOpaqueCache.xml')==b'<opaque>Preserve cache bytes</opaque>'
        assert z.read('customXml/syntheticExtension.xml')==b'<vendor-extension>Keep untouched</vendor-extension>'

def test_synthetic_b001_42_lines_1324_units(client,admin):
    # Structural/count test only. These are not the absent original B001 assets or prices.
    from .test_pi_api import B001_HEADER
    wb=Workbook();ws=wb.active;ws.title='Asset List';ws.append(B001_HEADER)
    for i in range(42):
        ws.append(['SYNTHETIC-SO','Local',f'SYNTHETIC-{i}','Dell Latitude 5430 i5-1245U','Synthetic site',
                   31 if i<41 else 53,'Laptop','Dell Latitude 5430','i5-1245U','12th Gen','16GB','256GB SSD','Synthetic structural test'])
    priced=io.BytesIO();wb.save(priced)
    # Import trusted historical test evidence once, then reuse without typing prices into this valuation.
    ws.cell(1,14,'Market Unit Price')
    for i in range(2,44):ws.cell(i,14,10000)
    refs=io.BytesIO();wb.save(refs)
    assert client.post('/api/pi/reference-packs',headers=admin,files={'file':('synthetic-history.xlsx',refs.getvalue())},
           data={'reference_date':'2026-10-26'}).status_code==200
    b=upload(client,admin,priced.getvalue()).json()
    assert b['lines']==42 and b['units']==1324 and b['qty_column']=='Net Leased Qty'
    j,data=finish(client,admin,b)
    assert j['workflow']['priced_lines']==42
    out=load_workbook(io.BytesIO(data));sheet=out['Asset List']
    assert sheet['N1'].value=='Market Unit Price' and sheet['O2'].value=='=IF(ISNUMBER(N2),N(F2)*N2,"")'
    assert {sheet.cell(i,14).value for i in range(2,44)} == {10000}
    assert client.get(f"/api/pi/batches/{b['id']}",headers=admin).json()['estimated_value']=='13240000.00'


def test_backoff_job_is_not_claimed_until_due(client,admin):
    from app.procurement.models import PIResearchJob,PIResearchTask,PIProduct
    from app.procurement import research
    from datetime import timedelta
    with SessionLocal() as db:
        p=PIProduct(product_key='test-backoff',category='other',label='Synthetic',spec={});db.add(p);db.flush()
        j=PIResearchJob(provider='offline',status='queued',total=1,params={});db.add(j);db.flush()
        db.add(PIResearchTask(job_id=j.id,product_id=p.id,next_attempt_at=datetime.now(timezone.utc)+timedelta(minutes=10)))
        db.commit();assert research.next_queued_job(db) is None


def test_similar_model_requires_explicit_factor(client,admin):
    from app.procurement.models import PIProduct,PIObservation
    data=sheet_bytes([['Model','Category','Grade','Qty'],['XEON GOLD 6130','PM-CPU','A',1]])
    b=upload(client,admin,data).json()
    with SessionLocal() as db:
        source=PIProduct(product_key='synthetic-comparable',category='server_cpu',label='Synthetic comparable CPU',spec={})
        db.add(source);db.flush()
        db.add(PIObservation(product_id=source.id,provider='manual',price_type='supplier_quote',match_quality='identical',
            price=Decimal('10000'),currency='INR',tax_basis='exclusive',observed_at=datetime.now(timezone.utc),status='verified',dedup_key='synthetic-comparable'))
        db.commit()
    j,_=finish(client,admin,b);assert j['workflow']['priced_lines']==0
    rules=client.get('/api/pi/rules',headers=admin).json()['active']['rules']
    rules['similar_model_factors']={'server_cpu|xeon-gold-6130':{'source_product_key':'synthetic-comparable','factor':'0.8','reason':'Synthetic approved adjustment test'}}
    assert client.post('/api/pi/rules',headers=admin,json={'rules':rules,'note':'Synthetic approved comparable factor'}).status_code==200
    j,_=finish(client,admin,b)
    a=client.get(f"/api/pi/batches/{b['id']}/assets",headers=admin).json()['items'][0]
    assert a['calculated_price']=='8000.00' and a['status']=='provisional' and a['research_methods']==['similar_model_estimate']


def test_automatic_workbook_download_requires_authentication(client,admin):
    b=upload(client,admin,b001_like_workbook()).json()
    j,_=finish(client,admin,b)
    assert client.get(f"/api/pi/jobs/{j['id']}/workbook").status_code==401
    assert client.post(f"/api/pi/batches/{b['id']}/analyze").status_code==401


def test_invalid_policy_rates_rejected():
    import pytest
    rules=default_rules().model_dump(mode='json')
    rules['categories']['laptop']['quantity_bands'][0]['discount_pct']='105'
    from app.procurement.rules import Rules
    with pytest.raises(ValueError):Rules.model_validate(rules)

def test_background_thread_finishes_without_per_job_commands(client,admin,monkeypatch):
    import time
    from fastapi.testclient import TestClient
    from app.main import create_app
    monkeypatch.setenv('CLPI_RESEARCH_WORKER','thread');get_settings.cache_clear()
    try:
        with TestClient(create_app()) as browser:
            b=upload(browser,admin,sheet_bytes([['Model','Category','Grade','Qty'],['XEON GOLD 6130','PM-CPU','A',1]])).json()
            response=browser.post(f"/api/pi/batches/{b['id']}/analyze",headers=admin)
            assert response.status_code==200,response.text
            deadline=time.monotonic()+15
            while time.monotonic()<deadline:
                job=browser.get(f"/api/pi/jobs/{response.json()['id']}",headers=admin).json()
                if job.get('workflow',{}).get('artifact'):break
                time.sleep(.1)
            assert job['workflow']['result']=='Research Unavailable',job
            assert browser.get(f"/api/pi/jobs/{job['id']}/workbook",headers=admin).status_code==200
    finally:
        worker.stop_in_process()
        if worker._thread:worker._thread.join(timeout=5)
        monkeypatch.setenv('CLPI_RESEARCH_WORKER','off');get_settings.cache_clear()

def test_cpu_classes_and_unknown_cosmetic_condition_are_not_collapsed():
    shared={'model':'Dell Latitude 5430','generation':'12th Gen','ram':'16GB','storage':'256GB','grade':'B'}
    a=identify({**shared,'cpu':'Core i5'})
    b=identify({**shared,'cpu':'Core i7'})
    assert a.product_key!=b.product_key
    c=identify({'model':'Dell Latitude 5430','condition':'working'})
    assert c.grade is None and 'condition_unknown' in c.risk_flags


@pytest.mark.parametrize('case,lines,units',[('B001',42,1324),('B002',2892,2892)])
def test_original_workbooks_queued_valuation_and_preservation(client,admin,monkeypatch,tmp_path,case,lines,units):
    original=B002.parent/f'CIN-20261026-{case}_Asset List.xlsx'
    if not original.exists():
        pytest.skip(f'Original {case} workbook not supplied')
    before=original.read_bytes()
    if case=='B002':
        history=tmp_path/'completed';history.mkdir();shutil.copy2(B002,history/B002.name)
        monkeypatch.setenv('CLPI_REFERENCE_DIRS',json.dumps([str(history)]));get_settings.cache_clear()
    b=upload(client,admin,before,name=original.name).json()
    assert b['layout']==case.lower() and b['lines']==lines and b['units']==units
    job,data=finish(client,admin,b)
    output=load_workbook(io.BytesIO(data)); source=load_workbook(io.BytesIO(before))
    asset=output[b['sheet_name']]; old=source[b['sheet_name']]
    if case=='B001':
        assert b['qty_column']=='Net Leased Qty' and asset.cell(b['header_row'],14).value=='Market Unit Price'
        assert job['workflow']['priced_lines']==0 and job['workflow']['result']=='Research Unavailable'
        for row in old:
            for cell in row:
                assert asset.cell(cell.row,cell.column).value==cell.value
    else:
        assert asset.cell(b['header_row'],10).value=='Market Unit Price'
        assert [c.value for c in asset[b['header_row']]].count('Market Unit Price')==1
        assert job['workflow']['priced_lines']==2891 and job['workflow']['result']=='Needs Review'
        for row in old:
            for cell in row:
                if cell.column!=10:
                    assert asset.cell(cell.row,cell.column).value==cell.value
        assert client.get(f"/api/pi/batches/{b['id']}",headers=admin).json()['estimated_value']=='13334100.00'
    allowed=set(job['workflow']['integrity']['fidelity']['modified_parts'])
    with ZipFile(io.BytesIO(before)) as aa,ZipFile(io.BytesIO(data)) as bb:
        assert set(aa.namelist())<=set(bb.namelist())
        for name in aa.namelist():
            if name not in allowed:
                assert aa.read(name)==bb.read(name)
    assert original.read_bytes()==before
    get_settings.cache_clear()
