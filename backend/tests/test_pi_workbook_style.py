"""Totals and presentation must survive the original-workbook preservation path."""
import io
from openpyxl import Workbook, load_workbook
from app.procurement.workbook_style import asset_style
from app.procurement.xlsx_patch import patch, cache_formula_values
from app.services.export import add_logo


def test_branded_total_filter_and_logo_preserve_source_values():
    source = Workbook()
    ws = source.active
    ws.title = "Assets"
    ws.append(["Serial", "Qty", "Market Unit Price", "Estimated Lot Value"])
    ws.append(["A-1", 3, None, None])
    ws.append(["A-2", 1, None, None])
    ws["A5"] = "Existing vendor note"
    stream = io.BytesIO(); source.save(stream)
    original = stream.getvalue()
    ws["C2"] = 125
    ws["D2"] = '=IF(ISNUMBER(C2),N(B2)*C2,"")'
    ws["D3"] = '=IF(ISNUMBER(C3),N(B3)*C3,"")'
    footer, refs = asset_style(ws, 1, 2, 3, 3, 4, 1)
    report = source.create_sheet("Valuation Summary")
    add_logo(report, "E1", 140)
    out = io.BytesIO(); source.save(out)
    data = cache_formula_values(out.getvalue(), {"Assets": {"D2": 375, **{r: 375 for r in refs}}})
    data, fidelity = patch(original, data, "Assets", 3, 4, 1, [2, 3], footer_rows=footer)
    result = load_workbook(io.BytesIO(data))
    styled = result["Assets"]
    assert styled["A5"].value == "Existing vendor note"
    assert styled["A2"].value == "A-1" and styled["B2"].value == 3
    assert styled["C3"].value is None
    assert styled.auto_filter.ref == "A1:D3"
    assert styled.freeze_panes == "B2"
    assert styled["A1"].fill.fgColor.rgb[-6:] == "0B2545"
    assert styled[refs[0]].value == "=SUM(D2:D3)"
    assert styled[refs[1]].value == "=SUBTOTAL(109,D2:D3)"
    assert load_workbook(io.BytesIO(data), data_only=True)["Assets"][refs[0]].value == 375
    assert len(result["Valuation Summary"]._images) == 1
    assert fidelity["mode"] == "minimal_ooxml_patch"
