"""End-to-end API tests: upload -> map -> review/correct -> commit -> value -> inspect -> export."""
import csv
import io
import zipfile
from datetime import date, timedelta

import httpx
from openpyxl import Workbook, load_workbook

from .conftest import login, make_user

TODAY = date.today()


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def csv_bytes(rows: list[list]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode()


def upload(client, h, record_type, name, content):
    r = client.post("/api/uploads", headers=h, data={"record_type": record_type},
                    files={"file": (name, content, "application/octet-stream")})
    return r


def stage(client, h, batch, sheets=None, defaults=None, save_as=None):
    names = sheets or [s["name"] for s in batch["sheets"]]
    sug = client.get(f"/api/uploads/{batch['id']}/suggest-mapping", headers=h,
                     params=[("sheets", n) for n in names]).json()
    r = client.post(f"/api/uploads/{batch['id']}/stage", headers=h,
                    json={"sheets": names, "mapping": sug["suggested"], "defaults": defaults or {},
                          "save_mapping_as": save_as})
    assert r.status_code == 200, r.text
    return r.json(), sug


INV_HEADER = ["SKU", "Brand", "Model", "Processor", "RAM", "Storage", "Grade", "Battery Health", "Faults",
              "Charger", "Warranty", "Qty", "Asking Price", "Tax Basis", "Refurb Cost", "Channel", "Target Margin",
              "Yield", "Location"]


def inv_row(sku, ram="16GB", refurb="800", qty=5, faults="none"):
    return [sku, "Dell", "Latitude 7490", "Intel Core i5-8350U", ram, "256GB SSD", "A", "86%", faults, "Yes",
            "6 months", qty, "20,000", "+GST", refurb, "retail", "15%", "95%", "Delhi"]


SALES_HEADER = ["Type", "Date", "Brand", "Model", "CPU", "RAM", "Storage", "Grade", "Battery", "Charger",
                "Warranty", "Price", "Tax", "Channel", "Customer", "Invoice"]


def sale_row(i, price, days_ago, ram="16 GB"):
    return ["sale", (TODAY - timedelta(days=days_ago)).strftime("%d/%m/%Y"), "Dell", "Latitude 7490",
            "i5-8350U", ram, "256 GB SSD", "Grade A", "88%", "yes", "6 months", f"₹{price:,}", "inclusive", "retail",
            f"Customer {i}", f"INV-{i}"]


def test_rbac_and_auth(client, admin):
    assert client.get("/api/dashboard").status_code == 401
    assert client.post("/api/auth/login", json={"email": "admin@test.local", "password": "wrong"}).status_code == 401
    sales = make_user(client, admin, "sales")
    proc = make_user(client, admin, "procurement")
    content = csv_bytes([INV_HEADER, inv_row("A1")])
    assert upload(client, sales, "inventory", "inv.csv", content).status_code == 403
    assert upload(client, proc, "inventory", "inv.csv", content).status_code == 200
    assert client.put("/api/settings/pricing", headers=proc, json={"config": {}}).status_code == 403
    assert client.get("/api/audit", headers=sales).status_code == 403
    assert client.post("/api/users", headers=proc, json={"email": "x@y.z", "name": "x", "role": "admin",
                                                          "password": "0123456789"}).status_code == 403
    # sales may add listings and record sales, not purchases
    r = client.post("/api/evidence/listings", headers=sales, json={"values": {
        "brand": "Dell", "model": "Latitude 7490", "price": "35000", "channel": "retail",
        "listing_date": TODAY.isoformat(), "cpu": "i5-8350U"}})
    assert r.status_code == 200, r.text
    r = client.post("/api/evidence/transactions", headers=sales, json={"values": {
        "kind": "purchase", "brand": "Dell", "model": "X", "price": "1", "transaction_date": TODAY.isoformat()}})
    assert r.status_code == 403


def test_unsafe_uploads_rejected(client, admin):
    assert upload(client, admin, "inventory", "a.xls", b"anything").status_code == 422
    assert upload(client, admin, "inventory", "a.xlsx", b"not a zip").status_code == 422
    assert upload(client, admin, "inventory", "a.csv", b"a,b\x00,c").status_code == 422
    assert upload(client, admin, "inventory", "a.csv", b"").status_code == 422
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/vbaProject.bin", "macro")
    r = upload(client, admin, "inventory", "a.xlsx", buf.getvalue())
    assert r.status_code == 422 and "macro" in r.json()["detail"]


def test_full_workflow(client, admin):
    # 1. Upload a two-sheet workbook of inventory.
    content = xlsx({
        "Lot A": [["Supplier stock list"], INV_HEADER, inv_row("A1"), inv_row("A2", ram="sixteen"),
                  inv_row("A1"), inv_row("A3", refurb="", qty=3)],
        "Lot B": [INV_HEADER, inv_row("B1", refurb="0", faults="")],
        "Notes": [["ignore me"]],
    })
    r = upload(client, admin, "inventory", "stock.xlsx", content)
    assert r.status_code == 200, r.text
    batch = r.json()
    assert [s["name"] for s in batch["sheets"]] == ["Lot A", "Lot B", "Notes"]
    # 2. Map columns (auto-suggested) across two selected sheets; the title row above Lot A's header is skipped.
    staged, sug = stage(client, admin, batch, sheets=["Lot A", "Lot B"], defaults={"lot_ref": "LOT-7"},
                        save_as="Supplier template")
    assert {"brand", "model", "cpu", "ram_gb", "storage_gb", "quantity", "asking_price"} <= set(sug["suggested"])
    assert staged["summary"]["total"] == 5 and staged["selected_sheets"] == ["Lot A", "Lot B"]
    rows = {(r["sheet"], r["row_number"]): r for r in client.get(f"/api/uploads/{batch['id']}/rows", headers=admin).json()}
    lot_b = rows[("Lot B", 2)]
    assert lot_b["status"] in ("valid", "warning")
    assert lot_b["normalized"]["refurbishment_cost"] == 0  # explicit zero preserved
    assert lot_b["normalized"]["faults"] is None  # blank = unknown, not "no faults"
    bad = rows[("Lot A", 4)]
    assert bad["status"] == "invalid" and "ram_gb" in [e["field"] for e in bad["errors"]]
    dup = rows[("Lot A", 5)]
    assert dup["status"] == "duplicate" and dup["duplicate_of"] == "row:Lot A:3"
    missing = rows[("Lot A", 6)]
    assert missing["normalized"]["refurbishment_cost"] is None  # blank = missing, not zero
    # 3. Downloadable import issues keep the original values.
    issues = client.get(f"/api/uploads/{batch['id']}/issues?format=csv", headers=admin)
    assert issues.status_code == 200 and "sixteen" in issues.text and "duplicate" in issues.text
    xl = client.get(f"/api/uploads/{batch['id']}/issues?format=xlsx", headers=admin)
    assert load_workbook(io.BytesIO(xl.content)).active.max_row >= 3
    # 4. Correct the invalid row; the correction is recorded against the original value.
    r = client.patch(f"/api/staging-rows/{bad['id']}", headers=admin,
                     json={"changes": {"ram_gb": "16"}, "reason": "Supplier confirmed 16 GB"})
    assert r.status_code == 200 and r.json()["status"] in ("valid", "warning")
    assert r.json()["corrections"] == {"ram_gb": "16"}
    assert r.json()["mapped_raw"]["ram_gb"] == "sixteen"
    # 5. Commit: valid rows only; the duplicate is not committed.
    r = client.post(f"/api/uploads/{batch['id']}/commit", headers=admin)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 4 and r.json()["not_committed"] == 1
    items = client.get("/api/inventory", headers=admin).json()
    a2 = next(i for i in items if i["item_ref"] == "A2")
    assert a2["raw_values"]["ram_gb"] == "sixteen" and a2["corrections"] == {"ram_gb": "16"} and a2["ram_gb"] == 16
    assert a2["lot_ref"] == "LOT-7"
    assert a2["asking_tax_basis"] == "exclusive" and a2["asking_price"] == 20000
    # 6. Import comparable evidence: completed sales (CSV), one stale, one other config.
    sales_rows = [SALES_HEADER] + [sale_row(i, 35000 + 500 * i, 5 + 7 * i) for i in range(6)] + \
                 [sale_row(90, 36000, 400), sale_row(91, 30000, 10, ram="8 GB")]
    b2 = upload(client, admin, "transaction", "sales.csv", csv_bytes(sales_rows)).json()
    staged2, _ = stage(client, admin, b2)
    assert staged2["summary"]["by_status"].get("invalid", 0) == 0, client.get(
        f"/api/uploads/{b2['id']}/rows?status=invalid", headers=admin).json()
    assert client.post(f"/api/uploads/{b2['id']}/commit", headers=admin).json()["created"] == 8
    # Re-uploading the same file flags every row as a duplicate of stored evidence.
    b3 = upload(client, admin, "transaction", "sales-again.csv", csv_bytes(sales_rows)).json()
    staged3, _ = stage(client, admin, b3)
    assert staged3["summary"]["by_status"] == {"duplicate": 8}
    # 7. Value the item.
    r = client.post("/api/valuations", headers=admin, json={"inventory_item_ids": [a2["id"]]})
    assert r.status_code == 200, r.text
    summary = r.json()[0]
    assert summary["status"] == "ok" and summary["selling_basis"] == "completed_sales"
    assert summary["target_buy"] is not None and summary["max_safe_buy"] is not None
    # 8. Inspect sources: included and excluded comparables with reasons.
    d = client.get(f"/api/valuations/{summary['id']}", headers=admin).json()
    comps = d["result"]["comparables"]
    assert sum(c["included"] for c in comps) == 6
    reasons = " | ".join(c["exclusion"] or "" for c in comps)
    assert "Stale" in reasons and "Materially incomparable: RAM (GB) 8 vs subject 16" in reasons
    assert all(c["source_document"].startswith("Upload #") for c in comps)
    assert d["result"]["confidence"]["note"].endswith("It is NOT a calibrated probability.")
    # 9. Reproduce and export.
    rep = client.post(f"/api/valuations/{summary['id']}/reproduce", headers=admin).json()
    assert rep["reproduced"] is True
    x = client.get(f"/api/valuations/{summary['id']}/export", headers=admin)
    wb = load_workbook(io.BytesIO(x.content))
    assert {"Summary", "Economics", "Comparables", "Assumptions", "Evidence quality"} <= set(wb.sheetnames)
    allx = client.get("/api/exports/valuations?format=csv", headers=admin)
    assert allx.status_code == 200 and "Max safe buy / unit" in allx.text
    # 10. Record an outcome and compare it to the prediction.
    r = client.post("/api/outcomes", headers=admin, json={"inventory_item_id": a2["id"], "kind": "sale",
                                                          "price": "36500", "transaction_date": TODAY.isoformat(),
                                                          "tax_basis": "inclusive", "channel": "retail"})
    assert r.status_code == 200, r.text
    out = client.get("/api/outcomes", headers=admin).json()[0]
    assert out["predicted"] == summary["selling_p50"] and out["difference"] is not None
    # Audit trail covers the flow.
    actions = {e["action"] for e in client.get("/api/audit", headers=admin).json()}
    assert {"import.uploaded", "import.staged", "import.row_corrected", "import.committed",
            "valuation.created", "export.valuation", "outcome.recorded"} <= actions


def test_manual_entry_validation_and_edit_history(client, admin):
    r = client.post("/api/inventory", headers=admin, json={"values": {"brand": "HP", "model": "EliteBook 840 G5"}})
    assert r.status_code == 422  # quantity required
    r = client.post("/api/inventory", headers=admin, json={"values": {
        "brand": "HP", "model": "EliteBook 840 G5", "quantity": "2", "cpu": "i5-8350U", "ram_gb": "8",
        "storage_gb": "256", "condition_grade": "B"}})
    assert r.status_code == 200
    iid = r.json()["id"]
    r = client.patch(f"/api/inventory/{iid}", headers=admin, json={"values": {"ram_gb": "16"}, "reason": "upgrade"})
    assert r.json()["changed"] == ["ram_gb"]
    hist = client.get(f"/api/inventory/{iid}", headers=admin).json()["edit_history"]
    assert hist[0]["previous"] == "8" and hist[0]["new"] == "16"


def test_insufficient_evidence_and_settings_versioning(client, admin):
    r = client.post("/api/inventory", headers=admin, json={"values": {
        "brand": "Apple", "model": "MacBook Air M1", "quantity": "1", "cpu": "Apple M1", "ram_gb": "8",
        "storage_gb": "256", "condition_grade": "A"}})
    run = client.post("/api/valuations", headers=admin, json={"inventory_item_ids": [r.json()["id"]]}).json()[0]
    assert run["status"] == "insufficient_evidence" and run["target_buy"] is None
    s = client.get("/api/settings/pricing", headers=admin).json()
    cfg = s["config"]
    cfg["default_target_margin_pct"] = 12
    cfg["adjustments"] = [{"id": "x", "dimension": "ram_gb", "comparable_value": "8", "subject_value": "16",
                           "method": "percent", "value": 7, "basis": "test"}]
    r = client.put("/api/settings/pricing", headers=admin, json={"config": cfg, "note": "set margin"})
    assert r.status_code == 200 and r.json()["version"] == s["version"] + 1
    cfg["adjustments"][0]["dimension"] = "colour"
    assert client.put("/api/settings/pricing", headers=admin, json={"config": cfg}).status_code == 422
    assert len(client.get("/api/settings/pricing/versions", headers=admin).json()) == 2


def test_demo_data_is_separated_and_removable(client, admin):
    r = client.post("/api/admin/demo", headers=admin)
    assert r.json()["loaded"] is True
    dash = client.get("/api/dashboard", headers=admin).json()
    assert dash["demo_data_present"] is True and dash["market_coverage"]["live"] is False
    assert dash["evidence"]["completed_sales"]["real"]["total"] == 0
    items = client.get("/api/inventory?synthetic=true", headers=admin).json()
    runs = client.post("/api/valuations", headers=admin, json={"inventory_item_ids": [i["id"] for i in items]}).json()
    statuses = {r["status"] for r in runs}
    assert {"ok", "insufficient_evidence", "needs_input"} <= statuses
    d = client.get(f"/api/valuations/{runs[0]['id']}", headers=admin).json()
    assert d["result"]["subject"]["is_synthetic"] is True
    assert all(c["is_synthetic"] for c in d["result"]["comparables"])
    bt = client.post("/api/validation/backtest", headers=admin, json={"synthetic": True}).json()
    assert bt["status"] == "ok" and bt["head_to_head"]["n"] > 0
    assert "p25_p75" in bt["interval_coverage"]
    real_bt = client.post("/api/validation/backtest", headers=admin, json={"synthetic": False}).json()
    assert real_bt["status"] == "insufficient_history"
    lots = client.get("/api/lots", headers=admin).json()
    assert lots[0]["is_synthetic"] and lots[0]["lines"] == 6
    removed = client.delete("/api/admin/demo", headers=admin).json()["removed"]
    assert removed["inventory"] == 6 and removed["valuations"] == len(runs)
    assert client.get("/api/dashboard", headers=admin).json()["demo_data_present"] is False


def test_rule_extraction_requires_review(client, admin):
    content = csv_bytes([["Brand", "Model", "Price", "Channel", "Listing date", "Title"],
                         ["Dell", "Latitude 7490", "34000", "retail", TODAY.isoformat(),
                          "Dell Latitude 7490 i5-8350U 16GB RAM 256GB SSD Grade A with charger"]])
    b = upload(client, admin, "listing", "l.csv", content).json()
    stage(client, admin, b)
    row = client.get(f"/api/uploads/{b['id']}/rows", headers=admin).json()[0]
    r = client.post(f"/api/staging-rows/{row['id']}/extract", headers=admin, json={"use_ai": False}).json()
    assert r["status"] == "needs_review" and r["suggestions"]["cpu"]["value"] == "i5-8350U"
    assert r["normalized"]["cpu"] is None  # nothing applied until accepted
    r = client.post(f"/api/staging-rows/{row['id']}/suggestions", headers=admin,
                    json={"accept": ["cpu", "ram_gb", "storage_gb", "condition_grade"], "reject": ["charger_included", "storage_type"]})
    rj = r.json()
    assert rj["normalized"]["cpu"] == "i5-8350U" and rj["normalized"]["ram_gb"] == 16 and rj["suggestions"] is None
    # AI is not configured: the endpoint says so instead of failing silently.
    assert client.post(f"/api/staging-rows/{row['id']}/extract", headers=admin,
                       json={"use_ai": True}).status_code == 503
    assert client.get("/api/ai/status", headers=admin).json()["configured"] is False


def test_connector_feed_requires_configuration_and_records_failures(client, admin):
    conns = client.get("/api/connectors", headers=admin).json()
    feed = next(c for c in conns if c["adapter"] == "http_json_feed")
    assert feed["status"] == "not_configured" and feed["config_problems"]
    assert client.post(f"/api/connectors/{feed['id']}/run", headers=admin).status_code == 502
    r = client.put(f"/api/connectors/{feed['id']}", headers=admin, json={"config": {"api_key": "secret"}})
    assert r.status_code == 422  # secrets never stored in the database
    cfg = {"url": "https://feed.example.test/listings", "terms_confirmed": True, "records_path": "data.items",
           "default_channel": "retail", "default_tax_basis": "inclusive", "platform": "Example feed",
           "field_map": {"brand": "brand", "model": "model", "cpu": "cpu", "ram_gb": "ram", "storage_gb": "ssd",
                         "condition_grade": "grade", "price": "price", "listing_date": "posted",
                         "source_url": "url", "source_record_id": "id", "seller": "seller"}}
    r = client.put(f"/api/connectors/{feed['id']}", headers=admin, json={"config": cfg, "enabled": True})
    assert r.status_code == 200, r.text

    from app.connectors.adapters import run_feed
    from app.db import SessionLocal
    from app.models import Connector, MarketListing
    from app.pricing.config import default_config

    def handler(request):
        return httpx.Response(200, json={"data": {"items": [
            {"id": "f1", "brand": "Dell", "model": "Latitude 7490", "cpu": "i5-8350U", "ram": "16GB", "ssd": "256GB",
             "grade": "A", "price": "35,500", "posted": TODAY.isoformat(), "url": "https://feed.example.test/f1",
             "seller": "S1"},
            {"id": "f2", "brand": "Dell", "model": "Latitude 7490", "price": "not a price", "posted": TODAY.isoformat()},
        ]}})

    with SessionLocal() as db:
        conn = db.get(Connector, feed["id"])
        res = run_feed(db, conn, default_config(), None, transport=httpx.MockTransport(handler))
        db.commit()
        assert res == {"created": 1, "rejected": 1, "duplicates_skipped": 0, "errors_sample": res["errors_sample"]}
        rec = db.query(MarketListing).one()
        assert rec.source_type == "connector" and rec.source_url == "https://feed.example.test/f1"
        res2 = run_feed(db, conn, default_config(), None, transport=httpx.MockTransport(handler))
        assert res2["duplicates_skipped"] == 1

        def failing(request):
            return httpx.Response(503)
        try:
            run_feed(db, conn, default_config(), None, transport=httpx.MockTransport(failing))
        except Exception:
            pass
        db.commit()
    feed = next(c for c in client.get("/api/connectors", headers=admin).json() if c["adapter"] == "http_json_feed")
    assert feed["status"] == "error" and feed["failure_count"] == 1 and feed["last_success_at"]
    assert "HTTP 503" in feed["last_error"]


def test_login_after_role_change(client, admin):
    proc = make_user(client, admin, "procurement")
    me = client.get("/api/auth/me", headers=proc).json()
    assert "import_inventory" in me["permissions"] and "edit_settings" not in me["permissions"]
    users = client.get("/api/users", headers=admin).json()
    uid = next(u["id"] for u in users if u["role"] == "procurement")
    client.patch(f"/api/users/{uid}", headers=admin, json={"is_active": False})
    assert client.get("/api/auth/me", headers=proc).status_code == 401
    assert login.__name__ == "login"


TEMPLATE_ROWS = {
    "inventory": {"Brand": "Dell", "Model": "Latitude 7490", "Exact CPU": "i5-8350U", "RAM (GB)": 16,
                  "Storage capacity (GB)": 256, "Storage type": "SSD", "Condition grade": "A", "Quantity": 4,
                  "Supplier asking price (per unit)": 21000, "Asking price tax basis": "exclusive",
                  "Charger included": "yes", "Faults": "none", "Intended channel": "retail"},
    "listing": {"Brand": "Dell", "Model": "Latitude 7490", "Exact CPU": "i5-8350U", "RAM (GB)": 16,
                "Storage capacity (GB)": 256, "Condition grade": "A", "Asking price (per unit)": 34500,
                "Channel": "retail", "Listing date": "TODAY", "Tax basis": "inclusive", "Seller": "Shop 1"},
    "quote": {"Brand": "Dell", "Model": "Latitude 7490", "Supplier": "Supplier 1", "Quoted unit price": 19000,
              "Quote date": "TODAY", "Tax basis": "exclusive"},
    "transaction": {"Brand": "Dell", "Model": "Latitude 7490", "Transaction type": "sale",
                    "Transaction date": "TODAY", "Realised unit price": 33000, "Channel": "retail",
                    "Tax basis": "inclusive"},
}


def test_downloadable_templates_round_trip(client, admin):
    """The sample Excel templates offered for download, once filled in, upload with every column recognised."""
    from datetime import datetime
    from pathlib import Path
    root = Path(__file__).resolve().parents[2]
    for rt, values in TEMPLATE_ROWS.items():
        path = root / "templates" / f"{rt}_template.xlsx"
        assert path.exists(), f"missing template {path.name}"
        wb = load_workbook(path)
        ws = wb["Data"]
        headers = [c.value for c in ws[1]]
        for k, v in values.items():
            ws.cell(2, headers.index(k) + 1, datetime.now() if v == "TODAY" else v)
        buf = io.BytesIO()
        wb.save(buf)
        r = upload(client, admin, rt, f"{rt}.xlsx", buf.getvalue())
        assert r.status_code == 200, r.text
        batch = r.json()
        data_sheet = next(s for s in batch["sheets"] if s["name"] == "Data")
        assert data_sheet["rows"] == 1  # the 2,000 pre-formatted empty rows are ignored
        staged, sug = stage(client, admin, batch, sheets=["Data"])
        assert len(sug["suggested"]) == len(headers), f"{rt}: not every template column auto-mapped"
        row = client.get(f"/api/uploads/{batch['id']}/rows", headers=admin).json()[0]
        assert row["errors"] == [], (rt, row["errors"])
        assert client.post(f"/api/uploads/{batch['id']}/commit", headers=admin).json()["created"] == 1
